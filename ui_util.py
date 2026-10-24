"""
ui_util.py - Tabelle modificabili che conservano le modifiche anche cambiando modulo.

st.data_editor tiene le modifiche in uno stato interno che Streamlit dimentica quando il widget non viene disegnato (per esempio
tornando all'Hub). Qui la tabella "vera" sta in st.session_state[chiave]; a ogni modifica una callback la aggiorna e cambia la
chiave del widget (nuova versione), cosi' il widget riparte pulito dai dati aggiornati.
"""
import pandas as pd
import streamlit as st

ss = st.session_state


def _applica_modifiche(df, stato):
    df = df.copy().reset_index(drop=True)
    for i, ch in (stato.get("edited_rows") or {}).items():
        for col, val in ch.items():
            if col in df.columns and int(i) < len(df):
                df[col] = df[col].astype(object)
                df.at[int(i), col] = val
    canc = [int(i) for i in (stato.get("deleted_rows") or []) if int(i) < len(df)]
    if canc:
        df = df.drop(df.index[canc]).reset_index(drop=True)
    agg = stato.get("added_rows") or []
    if agg:
        df = pd.concat([df, pd.DataFrame(agg, columns=df.columns)], ignore_index=True)
    return df


def esempi_attivi():
    """I dati di esempio nelle tabelle sono spenti: la suite parte vuota. Per riaccenderli (prove, demo) imposta ESEMPI_ATTIVI=1."""
    import os
    return os.environ.get("ESEMPI_ATTIVI", "").strip().lower() in ("1", "true", "si", "sì", "yes")


def senza_esempi(f):
    """Decoratore di `_esempio()`: quando gli esempi sono spenti restituisce le stesse tabelle, ma vuote (colonne e tipi restano). `f.originale()` da' comunque l'esempio (serve ai modelli Excel)."""
    import functools
    import pandas as pd

    def _vuoti(x):
        if isinstance(x, pd.DataFrame):
            return x.iloc[0:0].copy()
        if isinstance(x, (tuple, list)):
            return type(x)(_vuoti(i) for i in x)
        return x

    @functools.wraps(f)
    def w(*a, **k):
        r = f(*a, **k)
        return r if esempi_attivi() else _vuoti(r)
    w.originale = f
    return w


def sola_lettura():
    """True per gli utenti con ruolo «lettore»: possono consultare e scaricare i dossier, non modificare ne' salvare."""
    return (ss.get("utente") or {}).get("ruolo") == "lettore"


def tabella_modificabile(chiave, df_iniziale, column_config=None, altezza=None, num_rows="dynamic"):
    """Mostra la tabella e restituisce il DataFrame aggiornato. `chiave` identifica i dati in session_state (usa un prefisso del modulo)."""
    if chiave not in ss:
        ss[chiave] = df_iniziale.copy()
    ver = ss.setdefault(chiave + "_v", 0)
    key_ed = f"ed_{chiave}_{ver}"                      # nessun prefisso del modulo: lo stato del widget non va ripristinato a mano

    def _cb():
        ss[chiave] = _applica_modifiche(ss[chiave], ss.get(key_ed) or {})
        ss[chiave + "_v"] = ver + 1

    kw = {"height": altezza} if altezza else {}
    st.data_editor(ss[chiave], key=key_ed, on_change=_cb, num_rows="fixed" if sola_lettura() else num_rows, hide_index=True, column_config=column_config or {},
                   width="stretch", disabled=sola_lettura(), **kw)
    return ss[chiave]


def sostituisci_tabella(chiave, df):
    """Sostituisce i dati (es. dopo una generazione) e riavvia l'editor con la nuova versione."""
    ss[chiave] = df.copy()
    ss[chiave + "_v"] = ss.get(chiave + "_v", 0) + 1


def salva_protetto(modulo, nome, payload, utente, audit):
    """Salva controllando che nessun altro abbia modificato lo stesso salvataggio dopo che lo hai aperto. Restituisce True se ha salvato."""
    import storage
    nome = (nome or "").strip()
    base = ss.get(f"_base_{modulo}") or {}
    atteso = base.get("agg", "") if base.get("nome") == nome else ""
    try:
        agg = storage.salva_pratica(nome, payload, autore=utente, modulo=modulo, atteso=atteso)
    except storage.ConflittoSalvataggio as c:
        ss[f"_conf_{modulo}"] = {"nome": nome, "payload": payload, "autore": c.autore, "agg": c.aggiornato, "aperto": bool(atteso)}
        return False
    ss.pop(f"_conf_{modulo}", None)
    ss[f"_base_{modulo}"] = {"nome": nome, "agg": agg}
    return True                                   # il registro attivita' scrive da solo cosa e' cambiato (storage.salva_pratica)


def conflitto_salvataggio(modulo, utente, audit):
    """Se l'ultimo salvataggio ha trovato una modifica altrui, mostra l'avviso e il pulsante per sovrascrivere comunque."""
    import storage
    c = ss.get(f"_conf_{modulo}")
    if not c:
        return
    chi = f"da {c['autore']} " if c["autore"] else ""
    if c["aperto"]:
        st.warning(f"⚠️ «{c['nome']}» è stato modificato {chi}alle {c['agg'][11:16]} (UTC) dopo che l'avevi aperto. Salvando ora sostituiresti il suo lavoro.")
    else:
        st.warning(f"⚠️ Esiste già un salvataggio «{c['nome']}» (ultima modifica {chi}alle {c['agg'][11:16]} UTC) che non hai aperto. Salvando lo sostituiresti.")
    a, b = st.columns(2)
    if a.button("Sovrascrivi comunque", key=f"sv_forza_{modulo}", width="stretch"):
        agg = storage.salva_pratica(c["nome"], c["payload"], autore=utente, modulo=modulo)
        ss[f"_base_{modulo}"] = {"nome": c["nome"], "agg": agg}
        ss.pop(f"_conf_{modulo}", None)
        audit(f"{modulo}_salvato_forzato", c["nome"])
        st.rerun()
    if b.button("Annulla", key=f"sv_annulla_{modulo}", width="stretch"):
        ss.pop(f"_conf_{modulo}", None)
        st.rerun()


def pannello_archivio(modulo, etichetta, raccogli, applica, audit=None, utente=""):
    """Riquadro 'Salva / apri' per un modulo. `raccogli()` -> payload (dict JSON); `applica(payload, nome)` ripristina i dati."""
    import storage
    audit = audit or (lambda *a, **k: None)
    with st.expander(f"💾 Salva e riapri {etichetta}"):
        if not storage.database_persistente():
            st.warning("Database esterno non configurato: i salvataggi si perdono al riavvio dell'app (vedi DATABASE_URL nei Secrets).")
        a, b = st.columns([3, 1])
        nome = a.text_input("Nome", key=f"sv_nome_{modulo}", placeholder="es. Portafoglio 2027")
        b.markdown("<div style='height:1.9rem'></div>", unsafe_allow_html=True)
        if sola_lettura():
            st.info("Il tuo ruolo è «lettore»: puoi consultare e aprire i salvataggi, ma non modificarli.")
        if b.button("💾 Salva", key=f"sv_salva_{modulo}", width="stretch", disabled=sola_lettura()):
            try:
                if salva_protetto(modulo, nome, raccogli(), utente, audit):
                    st.success("Salvato.")
            except Exception as e:
                st.error(f"Impossibile salvare: {e}")
        conflitto_salvataggio(modulo, utente, audit)
        try:
            elenco = storage.elenca_pratiche(modulo=modulo)
            if storage.elenca_pratiche(modulo=modulo, demo=True) and st.checkbox("Mostra anche i dati dimostrativi", key=f"sv_demo_{modulo}", help="Salvataggi di prova «DEMO - …»: non toccano i dati veri."):
                elenco = storage.elenca_pratiche(modulo=modulo, demo=True) + elenco
        except Exception as e:
            elenco = []
            st.error(f"Database non disponibile: {e}")
        if elenco:
            nomi = [p["nome"] for p in elenco]
            scelta = st.selectbox("Salvati", nomi, key=f"sv_scelta_{modulo}",
                                  format_func=lambda n: next(f"{p['nome']}  ({p['aggiornato'][:16]}" + (f", {p['autore']}" if p["autore"] else "") + ")" for p in elenco if p["nome"] == n))
            c1, c2 = st.columns(2)
            if c1.button("📂 Apri", key=f"sv_apri_{modulo}", width="stretch"):
                payload = storage.carica_pratica(scelta, modulo=modulo)
                if payload:
                    audit(f"{modulo}_aperto", scelta)
                    meta = storage.meta_pratica(scelta, modulo=modulo)
                    ss[f"_base_{modulo}"] = {"nome": scelta, "agg": meta["aggiornato"] if meta else ""}
                    ss.pop(f"_conf_{modulo}", None)
                    applica(payload, scelta)
                    st.rerun()
            if c2.button("🗑️ Elimina", key=f"sv_elimina_{modulo}", width="stretch", disabled=sola_lettura()):
                storage.elimina_pratica(scelta, modulo=modulo)
                audit(f"{modulo}_eliminato", scelta)
                st.rerun()
            try:
                versioni = storage.storico_pratica(scelta, modulo=modulo)
            except Exception:
                versioni = []
            if versioni:
                with st.expander(f"🕘 Versioni precedenti di «{scelta}» ({len(versioni)})"):
                    vid = st.selectbox("Versione", [v["id"] for v in versioni], key=f"sv_ver_{modulo}",
                                       format_func=lambda i: next(f"{v['salvato'][:16]}" + (f" · {v['autore']}" if v["autore"] else "") for v in versioni if v["id"] == i))
                    if st.button("↩️ Ripristina questa versione", key=f"sv_rip_{modulo}", width="stretch"):
                        payload = storage.carica_versione(vid)
                        if payload:
                            audit(f"{modulo}_versione_ripristinata", f"{scelta} #{vid}")
                            meta = storage.meta_pratica(scelta, modulo=modulo)
                            ss[f"_base_{modulo}"] = {"nome": scelta, "agg": meta["aggiornato"] if meta else ""}
                            applica(payload, scelta)
                            st.rerun()
                    st.caption("Il ripristino apre i dati di quella versione; per renderla definitiva salvala di nuovo (la versione attuale finisce nello storico).")
        else:
            st.caption("Nessun salvataggio per questo modulo.")


def modello_excel(colonne, etichette, esempi):
    """Excel di esempio da scaricare, con intestazioni e qualche riga, per preparare un import."""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Modello"
    for j, c in enumerate(colonne, start=1):
        cell = ws.cell(row=1, column=j, value=etichette.get(c, c))
        cell.font, cell.fill = Font(color="FFFFFF", bold=True), PatternFill("solid", fgColor="1E293B")
        ws.column_dimensions[cell.column_letter].width = max(14, len(str(etichette.get(c, c))) + 4)
    for i, r in enumerate(esempi, start=2):
        for j, c in enumerate(colonne, start=1):
            ws.cell(row=i, column=j, value=r.get(c))
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


# ----------------------------------------------------------------------------
# IMPORT CON MAPPATURA DELLE COLONNE (export del gestionale / del commercialista)
# ----------------------------------------------------------------------------
def _norm(s):
    import re
    return re.sub(r"[^a-z0-9]", "", str(s).lower())


def indovina_mappa(colonne_file, campi, etichette, sinonimi=None):
    """Per ogni campo dell'app cerca la colonna del file che gli somiglia (nome, etichetta o sinonimi). Restituisce {campo: colonna|None}."""
    sinonimi = sinonimi or {}
    per_norm = {}
    for c in colonne_file:
        per_norm.setdefault(_norm(c), c)
    usate, mappa = set(), {}
    for campo in campi:
        cand = [campo, etichette.get(campo, "")] + list(sinonimi.get(campo, ()))
        trovata = None
        for x in cand:
            n = _norm(x)
            if n and n in per_norm and per_norm[n] not in usate:
                trovata = per_norm[n]
                break
        if trovata is None:                                    # corrispondenza parziale (es. "Data scadenza" ~ "Scadenza")
            for x in cand:
                n = _norm(x)
                for cn, c in per_norm.items():
                    if n and len(n) >= 4 and n in cn and c not in usate:
                        trovata = c
                        break
                if trovata:
                    break
        if trovata is not None:
            usate.add(trovata)
        mappa[campo] = trovata
    return mappa


def applica_mappa(df, mappa, campi):
    """DataFrame con le colonne dell'app, prese dal file secondo la mappa. Date in ISO, celle vuote a None."""
    import datetime as dt
    import pandas as pd
    righe = []
    for _, r in df.iterrows():
        riga = {}
        for c in campi:
            v = r[mappa[c]] if mappa.get(c) in df.columns else None
            if hasattr(v, "to_pydatetime"):
                v = v.to_pydatetime().date().isoformat()
            elif isinstance(v, (dt.datetime, dt.date)):
                v = (v.date() if isinstance(v, dt.datetime) else v).isoformat()
            elif v is not None and not isinstance(v, str) and pd.isna(v):
                v = None
            riga[c] = v
        righe.append(riga)
    out = pd.DataFrame(righe, columns=list(campi))
    return out.dropna(how="all").reset_index(drop=True).astype(object).where(lambda d: d.notna(), None)


def importa_dati(chiave, campi, etichette, sigla, idx, audit=None, sinonimi=None, dopo=None, trasforma=None):
    """Caricamento di un Excel/CSV con scelta guidata delle colonne. `dopo()` viene chiamata dopo un import riuscito."""
    import pandas as pd
    audit = audit or (lambda *a, **k: None)
    if sola_lettura():
        st.caption("Importazione non disponibile per il ruolo «lettore».")
        return
    up = st.file_uploader("Importa da Excel/CSV", type=["xlsx", "csv"], key=f"up_{sigla}{idx}")
    if up is None:
        return
    try:
        d = pd.read_csv(up, sep=None, engine="python") if up.name.lower().endswith(".csv") else pd.read_excel(up)
    except Exception as e:
        st.error(f"File non leggibile: {e}")
        return
    if d.empty:
        st.error("Il file non contiene righe.")
        return
    guess = indovina_mappa(list(d.columns), campi, etichette, sinonimi)
    st.caption("Controlla a quale colonna del tuo file corrisponde ogni campo (la scelta è già proposta). Lascia «—» per i campi che non hai.")
    opz = ["—"] + [str(c) for c in d.columns]
    cols = st.columns(3)
    scelta = {}
    for i, c in enumerate(campi):
        pre = guess.get(c)
        scelta[c] = cols[i % 3].selectbox(etichette.get(c, c), opz, index=opz.index(str(pre)) if pre is not None and str(pre) in opz else 0, key=f"map_{sigla}{idx}_{_norm(up.name)}_{c}")
    mappa = {c: (next((x for x in d.columns if str(x) == v), None) if v != "—" else None) for c, v in scelta.items()}
    modo = st.radio("Dati già presenti", ["Sostituisci", "Aggiungi in coda"], horizontal=True, key=f"mod_{sigla}{idx}")
    if st.button("Importa il file", key=f"imp_{sigla}{idx}"):
        nuovo = applica_mappa(d, mappa, campi)
        if trasforma is not None:
            nuovo = trasforma(nuovo)
        if nuovo.empty:
            st.error("Nessuna riga valida: controlla la corrispondenza delle colonne.")
        else:
            att = ss.get(chiave)
            if modo == "Aggiungi in coda" and att is not None and len(att):
                nuovo = pd.concat([att, nuovo], ignore_index=True)
            sostituisci_tabella(chiave, nuovo)
            audit(f"{sigla}_import", up.name)
            if dopo:
                dopo()
            st.rerun()


def blocco_tabella(chiave, titolo, nota, campi, etichette, cfg, esempio_df, sigla, idx, audit=None, altezza=240, sinonimi=None):
    """Tabella modificabile + modello Excel + import guidato, con lo stesso aspetto in tutti i moduli."""
    st.subheader(titolo)
    if nota:
        st.caption(nota)
    tabella_modificabile(chiave, esempio_df, cfg, altezza=altezza)
    with st.expander(f"Importa da Excel/CSV: {titolo}"):
        st.download_button("📥 Modello Excel", modello_excel(campi, etichette, esempio_df.head(3).to_dict("records")), file_name=f"modello_{sigla}_{idx}.xlsx",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key=f"dl_{sigla}{idx}")
        importa_dati(chiave, campi, etichette, sigla, idx, audit, sinonimi)


def banner_esempio(flag, tabelle, key):
    """Avviso «dati di esempio» con pulsante per svuotare le tabelle. `tabelle`: [(chiave, colonne)]."""
    import pandas as pd
    if ss.get(flag):
        st.info("Dati di esempio con date relative a oggi. Sostituiscili con i tuoi: modifica le tabelle o importa un file.")
        if st.button("Svuota l'esempio e parti da zero", key=key, disabled=sola_lettura()):
            for chiave, cols in tabelle:
                sostituisci_tabella(chiave, pd.DataFrame(columns=cols))
            ss[flag] = False
            st.rerun()


def scarica_dossier(sigla, genera, nome_file, audit, titolo_nome):
    """Blocco «Genera dossier PDF + Excel». `genera()` -> (pdf, xlsx, id). Lo stato sta in session_state['rep_<sigla>']."""
    nome = st.text_input("Nome del documento", titolo_nome, key=f"{sigla}_nome")
    if st.button("📄 Genera dossier PDF ed Excel", type="primary", key=f"gen_{sigla}"):
        with st.spinner("Preparo i documenti..."):
            pdf, xlsx, rid = genera(nome)
            ss[f"rep_{sigla}"] = {"pdf": pdf, "xlsx": xlsx, "id": rid}
            audit(f"{sigla}_report", rid)
    rep = ss.get(f"rep_{sigla}")
    if rep:
        c1, c2 = st.columns(2)
        c1.download_button("📥 Scarica il dossier PDF", rep["pdf"], file_name=f"{nome_file}_{rep['id']}.pdf", mime="application/pdf", width="stretch", key=f"dlp_{sigla}")
        c2.download_button("📥 Scarica l'Excel con formule", rep["xlsx"], file_name=f"{nome_file}_{rep['id']}.xlsx",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", width="stretch", key=f"dlx_{sigla}")
