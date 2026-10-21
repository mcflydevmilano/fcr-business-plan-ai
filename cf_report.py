"""cf_report.py - Dossier PDF ed Excel (con formule) del modulo Cash Flow di cantiere."""
import io
import time

import charts
import core
from core import DISCLAIMER, fmt_num
from pdf_report import BLU, CHIARO, GIALLO_B, GRIGIO, ROSSO_B, VERDE_B, _PDF, _strumenti
from report_util import copertina

MAX_CANTIERI_PDF = 30


def costruisci_ctx_cf(out, stress, nome=""):
    firma = {"p": out["params"], "k": out["kpi"], "nome": nome}
    return {
        "cf": out, "stress": stress, "nome": nome or "", "modello": "nessuno", "pratica": nome or "",
        "report_id": core.hash_payload(firma, 10).upper(), "generato_il": time.strftime("%d/%m/%Y %H:%M"),
        "header_titolo": "FCR GENERAL CONTRACTOR - Cash Flow di cantiere",
        "header_sub": "Fabbisogno di cassa del portafoglio cantieri | " + (nome or "portafoglio"),
        "riga_cover": ("Portafoglio: " + nome) if nome else "Portafoglio cantieri",
        "titolo_documento": "FCR General Contractor - Cash Flow di cantiere",
        "footer_dx": "Powered by Aetheris Labs",
        "nota_cover": "Proiezione basata su ipotesi dell'utente (IVA esclusa): non costituisce garanzia di risultato ne' consulenza finanziaria.",
    }


def grafici_cf(out, stress):
    k = out["kpi"]
    fido = out["params"]["fido"] or None
    return {
        "mensile": charts.fig_png(charts.fig_cf_mensile(out["mesi"], fido, "light")),
        "crediti": charts.fig_png(charts.fig_cf_crediti(out["mesi"], "light")),
        "stress": charts.fig_png(charts.fig_cf_stress(stress, fido, "light")) if stress else None,
    }


def genera_pdf_cf(ctx, grafici=None):
    out, stress = ctx["cf"], ctx.get("stress") or {}
    q, k, mesi = out["params"], out["kpi"], out["mesi"]
    grafici = grafici if grafici is not None else grafici_cf(out, stress)
    pdf = _PDF(ctx)
    T = _strumenti(pdf, grafici)
    sezione, paragrafo, tabella, tile, immagine, spazio = T.sezione, T.paragrafo, T.tabella, T.tile, T.immagine, T.spazio
    n = [0]

    def prossima():
        n[0] += 1
        return n[0]

    if k["primo_mese_fuori_fido"]:
        esito = ("SCOPERTO OLTRE IL FIDO", ROSSO_B, f"Dal mese {k['primo_mese_fuori_fido']}: picco {fmt_num(k['picco_fabbisogno'])} EUR contro un fido di {fmt_num(q['fido'])} EUR")
    elif k["picco_fabbisogno"] > 0:
        esito = ("SCOPERTO DENTRO IL FIDO", GIALLO_B, f"Picco di {fmt_num(k['picco_fabbisogno'])} EUR al mese {k['mese_picco']} su un fido di {fmt_num(q['fido'])} EUR")
    else:
        esito = ("NESSUNO SCOPERTO", VERDE_B, f"La cassa resta positiva per tutti i {k['mesi']} mesi (minimo {fmt_num(k['cassa_minima'])} EUR)")
    copertina(pdf, T.tile, ctx, "Cash Flow di cantiere", "Dossier del cash flow", "Fabbisogno di cassa del portafoglio", [
        ("DATA", ctx["generato_il"]), ("ID REPORT", ctx["report_id"]), ("CANTIERI", str(k["n_cantieri"])), ("ORIZZONTE", f"{k['mesi']} mesi")],
        [("Contratti", fmt_num(k["contratti"]) + " EUR", f"{k['n_cantieri']} cantieri", BLU),
         ("Margine lordo", fmt_num(k["margine_lordo"]) + " EUR", f"{fmt_num(k['margine_lordo_pct'], 1)}% dei contratti", VERDE_B),
         ("Picco fabbisogno", fmt_num(k["picco_fabbisogno"]) + " EUR", f"mese {k['mese_picco']}" if k["mese_picco"] else "nessuno scoperto", GIALLO_B if k["picco_fabbisogno"] else VERDE_B),
         ("Cassa finale", fmt_num(k["cassa_finale"]) + " EUR", f"dopo {k['mesi']} mesi", BLU)], esito)

    pdf.add_page()
    sezione("Ipotesi e portafoglio", prossima(), 80)
    tabella(["Ipotesi", "Valore"], [
        ("Contratti totali (IVA esclusa)", fmt_num(k["contratti"]) + " EUR"), ("Costi diretti del portafoglio", fmt_num(k["costi_diretti"]) + " EUR"),
        ("Margine lordo", f"{fmt_num(k['margine_lordo'])} EUR ({fmt_num(k['margine_lordo_pct'], 1)}%)"),
        ("Costi generali mensili", fmt_num(q["costi_generali"]) + " EUR"), ("Giorni medi di pagamento dei fornitori", fmt_num(q["gg_pagamento"], 0)),
        ("Giorni medi di incasso (ponderati sui contratti)", fmt_num(k["gg_incasso_medio"], 0)),
        ("Cassa iniziale", fmt_num(q["cassa0"]) + " EUR"), ("Fido disponibile", fmt_num(q["fido"]) + " EUR"),
        ("Tasso sullo scoperto", fmt_num(q["tasso_scoperto"], 2) + "%"),
    ], [126, 60], ["L", "R"])

    sezione("Cash flow mese per mese", prossima(), 100)
    righe = [(str(r["Mese"]), str(r["Cantieri attivi"]), fmt_num(r["Incassi"]), fmt_num(-r["Pagamenti"]), fmt_num(-r["Costi generali"]),
              fmt_num(-r["Interessi"]), fmt_num(r["Flusso netto"]), fmt_num(r["Cassa finale"])) for r in mesi]
    righe.append(("Totale", "", fmt_num(k["incassi_tot"]), fmt_num(-k["pagamenti_tot"]), fmt_num(-k["costi_generali_tot"]),
                  fmt_num(-k["interessi_totali"]), fmt_num(sum(r["Flusso netto"] for r in mesi)), fmt_num(k["cassa_finale"])))
    tabella(["Mese", "Attivi", "Incassi", "Pagamenti", "Costi gen.", "Interessi", "Flusso", "Cassa"], righe,
            [14, 14, 26, 26, 26, 22, 28, 30], ["C", "R", "R", "R", "R", "R", "R", "R"], size=7.8, h=5.6)
    immagine("mensile", 160)
    immagine("crediti", 150)

    if stress:
        sezione("Stress test", prossima(), 90)
        paragrafo("Ciascun caso applica un peggioramento alle ipotesi base (incassi piu' lenti, costi maggiori, partenze ritardate).", 9)
        conf_ = [(c["Caso"][:34], fmt_num(c["Picco fabbisogno"]), (str(c["Mese del picco"]) if c["Mese del picco"] else "-"), fmt_num(c["Cassa minima"]),
                  fmt_num(c["Interessi"]), (str(c["Fuori fido dal mese"]) if c["Fuori fido dal mese"] else "no")) for c in _confronto(stress)]
        tabella(["Caso", "Picco", "Mese", "Cassa min.", "Interessi", "Fuori fido"], conf_, [66, 28, 14, 30, 24, 24], ["L", "R", "C", "R", "R", "C"], size=8)
        immagine("stress", 160)

    sezione("Cantieri del portafoglio", prossima(), 70)
    cant = sorted(out["cantieri"], key=lambda d: -d["Contratto"])
    righe = [(d["Cantiere"][:28], fmt_num(d["Contratto"]), f"{d['Inizio']}-{d['Fine']}", fmt_num(d["Margine"]), fmt_num(d["Margine %"], 1) + "%")
             for d in cant[:MAX_CANTIERI_PDF]]
    if len(cant) > MAX_CANTIERI_PDF:
        resto = cant[MAX_CANTIERI_PDF:]
        righe.append((f"Altri {len(resto)} cantieri", fmt_num(sum(d["Contratto"] for d in resto)), "", fmt_num(sum(d["Margine"] for d in resto)), ""))
    tabella(["Cantiere", "Contratto", "Mesi", "Margine", "Margine %"], righe, [62, 34, 26, 34, 30], ["L", "R", "C", "R", "R"], size=8, h=5.8)

    sezione("Punti di attenzione", prossima(), 40)
    if out["avvisi"]:
        for a in out["avvisi"]:
            paragrafo("- " + a, 9)
    else:
        paragrafo("I controlli automatici non hanno rilevato criticita' (cassa sempre positiva, nessun cantiere in perdita).", 9)

    sezione("Metodologia e limiti", prossima(), 70)
    paragrafo("Per ogni cantiere: anticipo alla partenza, saldo a fine lavori e stati di avanzamento mensili uguali per il resto; ogni fatturazione "
              "si incassa dopo i giorni indicati (le frazioni di mese sono ripartite tra due mesi); i costi sono uniformi sulla durata e si pagano dopo "
              "i giorni medi dei fornitori. Cassa(t) = cassa(t-1) + incassi - pagamenti - costi generali - interessi; gli interessi sono calcolati "
              "sullo scoperto del mese precedente.", 8.5)
    paragrafo("Limiti: importi al netto di IVA (da gestire a parte), avanzamento lineare dei lavori, nessun anticipo ai fornitori, trattenute di "
              "garanzia, ritenute, varianti, riserve, imposte o dividendi. I risultati dipendono dalle ipotesi inserite.", 8.5)
    spazio(24)
    y0 = pdf.get_y() + 2
    pdf.set_fill_color(*CHIARO)
    pdf.set_font("Helvetica", "I", 8)
    pdf.set_text_color(*GRIGIO)
    pdf.set_xy(14, y0)
    pdf.multi_cell(182, 4.1, "Avvertenza: " + DISCLAIMER.replace("modello di scoring interno", "modello di previsione").encode("latin-1", "replace").decode("latin-1"))
    return bytes(pdf.output())


def _confronto(stress):
    import cf
    return cf.confronto_stress(stress)


def genera_excel_cf(ctx):
    """Excel con formule: Input (celle gialle), Cantieri, Mensile (flusso, interessi e cassa ricalcolati dalle formule)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    out = ctx["cf"]
    q, mesi, k = out["params"], out["mesi"], out["kpi"]
    wb = Workbook()
    giallo, testata = PatternFill("solid", fgColor="FFF2CC"), PatternFill("solid", fgColor="1E293B")
    f_h, f_in = Font(color="FFFFFF", bold=True), Font(color="1D4ED8", bold=True)
    EUR = '#,##0;[Red]-#,##0'

    ws = wb.active
    ws.title = "Input"
    ws["A1"] = "FCR General Contractor - Cash Flow di cantiere (celle gialle modificabili)"
    ws["A1"].font = Font(bold=True, size=13)
    voci = [("Cassa iniziale (EUR)", q["cassa0"], EUR), ("Fido disponibile (EUR)", q["fido"], EUR), ("Tasso sullo scoperto (% annuo)", q["tasso_scoperto"], "0.00"),
            ("Costi generali mensili (EUR)", q["costi_generali"], EUR)]
    for i, (et, v, fmt) in enumerate(voci, start=3):
        ws.cell(row=i, column=1, value=et)
        c = ws.cell(row=i, column=2, value=v)
        c.fill, c.font, c.number_format = giallo, f_in, fmt
    ws["A8"] = "Nota: incassi e pagamenti dei cantieri sono calcolati dall'app (ritardi, SAL, anticipi); modificando le celle gialle si ricalcolano costi generali, interessi, cassa e fido."
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 18

    wc = wb.create_sheet("Cantieri")
    intest = ["Cantiere", "Contratto", "Inizio", "Fine", "Margine", "Margine %"]
    for j, t in enumerate(intest, start=1):
        c = wc.cell(row=1, column=j, value=t)
        c.fill, c.font = testata, f_h
    for i, d in enumerate(out["cantieri"], start=2):
        wc.cell(row=i, column=1, value=d["Cantiere"])
        wc.cell(row=i, column=2, value=d["Contratto"]).number_format = EUR
        wc.cell(row=i, column=3, value=d["Inizio"])
        wc.cell(row=i, column=4, value=d["Fine"])
        wc.cell(row=i, column=5, value=f"=B{i}*F{i}/100").number_format = EUR
        wc.cell(row=i, column=6, value=round(d["Margine %"], 4))
    ultimo = len(out["cantieri"]) + 1
    wc.cell(row=ultimo + 1, column=1, value="Totale").font = Font(bold=True)
    wc.cell(row=ultimo + 1, column=2, value=f"=SUM(B2:B{ultimo})").number_format = EUR
    wc.cell(row=ultimo + 1, column=5, value=f"=SUM(E2:E{ultimo})").number_format = EUR
    wc.column_dimensions["A"].width = 30
    for col in "BCDEF":
        wc.column_dimensions[col].width = 16

    wm = wb.create_sheet("Mensile")
    intest = ["Mese", "Cantieri attivi", "Incassi", "Pagamenti", "Costi generali", "Interessi", "Flusso netto", "Cassa finale", "Utilizzo scoperto", "Fido residuo"]
    for j, t in enumerate(intest, start=1):
        c = wm.cell(row=1, column=j, value=t)
        c.fill, c.font = testata, f_h
    for i, r in enumerate(mesi, start=2):
        prec = "Input!$B$3" if i == 2 else f"H{i - 1}"
        wm.cell(row=i, column=1, value=r["Mese"])
        wm.cell(row=i, column=2, value=r["Cantieri attivi"])
        wm.cell(row=i, column=3, value=round(r["Incassi"], 2)).number_format = EUR
        wm.cell(row=i, column=4, value=round(r["Pagamenti"], 2)).number_format = EUR
        wm.cell(row=i, column=5, value="=Input!$B$6").number_format = EUR
        wm.cell(row=i, column=6, value=f"=MAX(0,-{prec})*Input!$B$5/1200").number_format = EUR
        wm.cell(row=i, column=7, value=f"=C{i}-D{i}-E{i}-F{i}").number_format = EUR
        wm.cell(row=i, column=8, value=f"={prec}+G{i}").number_format = EUR
        wm.cell(row=i, column=9, value=f"=MAX(0,-H{i})").number_format = EUR
        wm.cell(row=i, column=10, value=f"=Input!$B$4-I{i}").number_format = EUR
    u = len(mesi) + 1
    wm.cell(row=u + 1, column=1, value="Totale").font = Font(bold=True)
    for col in "CDEFG":
        wm[f"{col}{u + 1}"] = f"=SUM({col}2:{col}{u})"
        wm[f"{col}{u + 1}"].number_format = EUR
    wm.cell(row=u + 3, column=1, value="Picco di scoperto").font = Font(bold=True)
    wm.cell(row=u + 3, column=3, value=f"=MAX(I2:I{u})").number_format = EUR
    wm.cell(row=u + 4, column=1, value="Cassa minima").font = Font(bold=True)
    wm.cell(row=u + 4, column=3, value=f"=MIN(H2:H{u})").number_format = EUR
    for col in "ABCDEFGHIJ":
        wm.column_dimensions[col].width = 17
    wm.freeze_panes = "B2"
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
