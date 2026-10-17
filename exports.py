"""
exports.py - Report PDF, Word (.docx) ed Excel con formule vive (.xlsx).

Tutti i generatori ricevono un dizionario `ctx` con:
  dati, scen, piano, rating, leve, piano_par, soglie, covenant, cov_res, stress,
  commesse_kpi, commesse_df, per_cliente, storico, scenari, dati_demo,
  report_id, modello, generato_il, pratica
"""
import io
import os

from core import (
    DISCLAIMER, SOGLIE_DEFAULT, fmt_eur, fmt_num, fmt_pct, fmt_x, ratio,
)
import charts

LOGO_PATH = "aetheris_labs_logo.jpg"


# ----------------------------------------------------------------------------
# GRAFICI PER I REPORT
# ----------------------------------------------------------------------------
def grafici_report(ctx):
    """Restituisce {nome: bytes PNG} (tema chiaro)."""
    s = ctx.get("soglie") or SOGLIE_DEFAULT
    cov = ctx.get("covenant") or {}
    piano, dati = ctx["piano"], ctx["dati"]
    out = {}
    base_ratio = ratio(ctx["scen"]["pfn"], ctx["scen"]["ebitda"])
    out["pfn_ebitda"] = charts.fig_png(
        charts.fig_pfn_ebitda(
            {"Scenario": piano}, s, cov.get("pfn_ebitda_max"), "light",
            anno_base=int(dati["anno_riferimento"]), base_val=base_ratio,
        )
    )
    out["dscr"] = charts.fig_png(charts.fig_dscr(piano, s, cov.get("dscr_min"), "light"))
    out["cascata"] = charts.fig_png(charts.fig_cascata(piano, ctx["scen"]["pfn"], "light"))
    out["storico"] = charts.fig_png(charts.fig_storico(ctx.get("storico") or [], dati, piano, "light"))
    if ctx.get("scenari") and len(ctx["scenari"]) > 1:
        out["confronto"] = charts.fig_png(charts.fig_confronto(ctx["scenari"], s, "light"))
    if ctx.get("per_cliente") is not None:
        out["commesse"] = charts.fig_png(charts.fig_commesse(ctx["per_cliente"], "light"))
    return out


def _valore_dettaglio(d):
    v = d["valore"]
    if d["chiave"] == "equity_ratio":
        return fmt_pct(100 * v)
    if v is None:
        return "n.d."
    return fmt_x(v)


def _soglie_testo(d):
    if d["chiave"] == "equity_ratio":
        return f">= {fmt_pct(100 * d['verde'], 0)} / {fmt_pct(100 * d['giallo'], 0)}"
    segno = "<" if d["verso"] == "min" else ">="
    return f"{segno} {fmt_num(d['verde'], 2)} / {fmt_num(d['giallo'], 2)}"


# ----------------------------------------------------------------------------
# PDF
# ----------------------------------------------------------------------------
def pdf_safe(text):
    """Sostituisce i caratteri non latin-1 (i font core di FPDF non li supportano)."""
    if text is None:
        return ""
    rep = {
        "€": "EUR", "–": "-", "—": "-", "’": "'", "‘": "'", "“": '"', "”": '"',
        "…": "...", "•": "-", "≥": ">=", "≤": "<=", "→": "->",
    }
    for k, v in rep.items():
        text = text.replace(k, v)
    return text.encode("latin-1", "replace").decode("latin-1")


def genera_pdf(ctx, relazione, grafici=None):
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    grafici = grafici if grafici is not None else grafici_report(ctx)

    class PDFReport(FPDF):
        def header(self):
            self.set_fill_color(15, 23, 42)
            self.rect(0, 0, 210, 26, "F")
            if os.path.exists(LOGO_PATH):
                try:
                    self.image(LOGO_PATH, 10, 4, 18)
                except Exception:
                    pass
            self.set_font("Helvetica", "B", 10)
            self.set_text_color(255, 255, 255)
            self.set_xy(32, 6)
            self.cell(0, 6, "FCR GENERAL CONTRACTOR - RISK MANAGEMENT REPORT", align="L")
            self.set_font("Helvetica", "", 8)
            self.set_xy(32, 13)
            self.cell(0, 6, "Powered by Aetheris Labs | Lead Eng: Ing. Lorenzo Di Ferdinando", align="L")
            self.set_y(32)

        def footer(self):
            self.set_y(-15)
            self.set_font("Helvetica", "I", 7)
            self.set_text_color(128, 128, 128)
            self.cell(
                0, 10,
                pdf_safe(f"Pagina {self.page_no()} | ID report {ctx['report_id']} | {ctx['generato_il']} | "
                         f"Modello AI: {ctx['modello']}"),
                align="C",
            )

    pdf = PDFReport()
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()
    nl = dict(new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    def riga(cells, widths, aligns, fill=False, h=6):
        for i, (txt, w, a) in enumerate(zip(cells, widths, aligns)):
            last = i == len(cells) - 1
            pdf.cell(w, h, pdf_safe(txt), border=1, align=a, fill=fill,
                     new_x=XPos.LMARGIN if last else XPos.RIGHT, new_y=YPos.NEXT if last else YPos.TOP)

    def titolo(txt):
        if pdf.get_y() > 255:
            pdf.add_page()
        pdf.set_font("Helvetica", "B", 11)
        pdf.set_text_color(15, 23, 42)
        pdf.cell(0, 7, pdf_safe(txt), **nl)

    def intestazione(cells, widths, aligns, size=8):
        pdf.set_font("Helvetica", "B", size)
        pdf.set_fill_color(30, 41, 59)
        pdf.set_text_color(255, 255, 255)
        riga(cells, widths, aligns, fill=True, h=7)
        pdf.set_font("Helvetica", "", size)
        pdf.set_text_color(0, 0, 0)

    def immagine(nome, w=175):
        png = grafici.get(nome)
        if not png:
            return
        try:
            from PIL import Image
            with Image.open(io.BytesIO(png)) as im:
                h = w * im.height / im.width
        except Exception:
            h = w * 0.5
        if pdf.get_y() + h > 280:
            pdf.add_page()
        pdf.image(io.BytesIO(png), x=(210 - w) / 2, w=w)
        pdf.ln(2)

    base, scen, rating, piano = ctx["dati"], ctx["scen"], ctx["rating"], ctx["piano"]

    pdf.set_font("Helvetica", "B", 15)
    pdf.set_text_color(15, 23, 42)
    pdf.cell(0, 8, "VALUTAZIONE DEL MERITO DI CREDITO E BUSINESS PLAN", **nl)
    pdf.set_font("Helvetica", "I", 10)
    pdf.set_text_color(100, 100, 100)
    pdf.cell(0, 6, "Analisi indicativa, scoring prospettico e parere per il Comitato Fidi", **nl)
    if ctx.get("pratica"):
        pdf.cell(0, 6, pdf_safe(f"Pratica: {ctx['pratica']}"), **nl)
    pdf.ln(2)

    if ctx.get("dati_demo"):
        pdf.set_font("Helvetica", "B", 9)
        pdf.set_text_color(185, 28, 28)
        pdf.multi_cell(0, 5, pdf_safe("ATTENZIONE: il report e' basato su dati di esempio non verificati."), **nl)
        pdf.ln(1)

    # 1. Sintesi
    titolo(f"1. Sintesi Indicatori (anno base {int(base['anno_riferimento'])}) - Base vs Scenario")
    w, al = [80, 50, 50], ["L", "R", "R"]
    intestazione(["  Parametro", "Base  ", "Scenario  "], w, al, 9)
    righe = [
        ("Ricavi delle Vendite", fmt_eur(base["ricavi"]), fmt_eur(scen["ricavi"])),
        ("EBITDA (MOL)", fmt_eur(base["ebitda"]), fmt_eur(scen["ebitda"])),
        ("Patrimonio Netto", fmt_eur(base["patrimonio_netto"]), fmt_eur(scen["pn"])),
        ("Posizione Finanziaria Netta", fmt_eur(base["pfn"]), fmt_eur(scen["pfn"])),
        ("PFN / EBITDA", fmt_x(ratio(base["pfn"], base["ebitda"])), fmt_x(rating["pfn_ebitda"])),
        ("PFN / Patrimonio Netto", fmt_x(ratio(base["pfn"], base["patrimonio_netto"])), fmt_x(rating["pfn_pn"])),
        ("Patrimonio Netto / Totale Attivo",
         fmt_pct(100 * base["patrimonio_netto"] / base["totale_attivo"]) if base["totale_attivo"] > 0 else "n.d.",
         fmt_pct(100 * rating["equity_ratio"])),
        ("DSCR (Anno 1 del piano)", "-", fmt_x(rating["dscr"]) if rating["dscr"] is not None else "n.d."),
        ("Punteggio scoring (max 8)", "-", str(rating["punti"])),
    ]
    fill = False
    for lab, a, b in righe:
        pdf.set_fill_color(*((241, 245, 249) if fill else (255, 255, 255)))
        riga(["  " + lab, a + "  ", b + "  "], w, al, fill=True)
        fill = not fill
    pdf.set_font("Helvetica", "B", 9)
    pdf.cell(0, 7, pdf_safe("Giudizio stimato (scenario): " + rating["classe"]), **nl)
    pdf.ln(2)

    # Dettaglio punteggio
    titolo("Dettaglio del punteggio (verde / giallo)")
    wd, ad = [70, 35, 50, 25], ["L", "R", "R", "R"]
    intestazione(["  Indicatore", "Valore  ", "Soglie  ", "Punti  "], wd, ad)
    for d in rating["dettaglio"]:
        riga(["  " + d["indicatore"], _valore_dettaglio(d) + "  ", _soglie_testo(d) + "  ", f"{d['punti']}/2  "], wd, ad)
    pdf.ln(3)

    # 2. Piano
    titolo("2. Proiezioni Economico-Finanziarie Pluriennali (partendo dallo scenario)")
    w2 = [34, 26, 26, 26, 26, 26, 26]
    a2 = ["L"] + ["R"] * 6
    intestazione(["  Anno", "Ricavi", "EBITDA", "FCF", "PFN", "PFN/EBITDA", "DSCR"], w2, a2)
    for r in piano:
        riga(["  " + r["Anno"], fmt_num(r["Ricavi"]), fmt_num(r["EBITDA"]), fmt_num(r["FCF"]), fmt_num(r["PFN"]),
              fmt_x(r["PFN/EBITDA"]), fmt_x(r["DSCR"]) if r["DSCR"] is not None else "n.d."], w2, a2)
    pdf.set_font("Helvetica", "I", 8)
    pdf.set_text_color(100, 100, 100)
    p = ctx["piano_par"]
    pdf.multi_cell(
        0, 4,
        pdf_safe(f"Importi in EUR. Ipotesi: crescita ricavi {fmt_num(p['g_ricavi'], 1)}%, EBITDA {fmt_num(p['g_ebitda'], 1)}%, "
                 f"capex {fmt_num(p['capex'])}, ammortamenti {fmt_num(p['amm'])}, rimborso annuo {fmt_num(p['rimborso'])}, "
                 f"tasso {fmt_num(p['tasso'], 2)}%, aliquota {fmt_num(p['aliquota'], 1)}%, CCN {fmt_num(p['ccn_pct'], 0)}% "
                 "della variazione dei ricavi. FCF = EBITDA - imposte - capex - interessi - variazione CCN."),
        **nl,
    )
    pdf.ln(2)

    # Grafici
    if pdf.get_y() + 100 > 280:
        pdf.add_page()
    titolo("Grafici")
    for nome in ("pfn_ebitda", "dscr", "cascata", "storico", "confronto"):
        immagine(nome)

    # Covenant
    if ctx.get("cov_res"):
        titolo("Verifica covenant")
        wc, ac = [38, 28, 28, 24, 24, 24, 24], ["L"] + ["R"] * 6
        intestazione(["  Anno", "PFN/EBITDA", "Limite", "Esito", "DSCR", "Limite", "Esito"], wc, ac)
        for c in ctx["cov_res"]:
            riga(["  " + c["Anno"], fmt_x(c["PFN/EBITDA"]), fmt_x(c["Limite PFN/EBITDA"]), c["Esito leva"],
                  fmt_x(c["DSCR"]) if c["DSCR"] is not None else "n.d.", fmt_x(c["Limite DSCR"]), c["Esito DSCR"]], wc, ac)
        pdf.ln(3)

    # Stress test
    if ctx.get("stress"):
        titolo("Stress test")
        ws, as_ = [76, 34, 28, 26, 26], ["L", "R", "R", "R", "R"]
        intestazione(["  Caso", "PFN/EBITDA a3", "DSCR min", "Punti", "Fascia"], ws, as_)
        for s in ctx["stress"]:
            riga(["  " + s["Caso"], fmt_x(s["PFN/EBITDA anno 3"]),
                  fmt_x(s["DSCR minimo"]) if s["DSCR minimo"] is not None else "n.d.", f"{s['Punteggio']}/8", s["Fascia"]], ws, as_)
        pdf.ln(3)

    # Commesse
    k = ctx.get("commesse_kpi")
    if k:
        titolo("Portafoglio commesse")
        wk, ak = [100, 80], ["L", "R"]
        intestazione(["  Indicatore", "Valore  "], wk, ak)
        voci = [
            ("Numero commesse", str(k["n_commesse"])),
            ("Importo totale contratti", fmt_eur(k["importo_totale"])),
            ("Backlog residuo", fmt_eur(k["backlog"])),
            ("Copertura ricavi (backlog / ricavi base)", fmt_x(k["copertura_ricavi"]) if k["copertura_ricavi"] is not None else "n.d."),
            ("Produzione eseguita non fatturata (negativo = anticipi)", fmt_eur(k["da_fatturare"])),
            ("Crediti verso clienti da commesse", fmt_eur(k["crediti"])),
            ("Margine medio ponderato sul backlog", fmt_pct(k["margine_backlog_pct"])),
            ("Concentrazione primo cliente / primi 3", f"{fmt_pct(k['top1_pct'], 0)} / {fmt_pct(k['top3_pct'], 0)}"),
        ]
        for a, b in voci:
            riga(["  " + a, b + "  "], wk, ak)
        pdf.ln(2)
        immagine("commesse")

    # Relazione
    titolo("3. Relazione e Parere del Risk Management")
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(50, 50, 50)
    pdf.multi_cell(0, 5, pdf_safe(relazione), **nl)
    pdf.ln(4)

    pdf.set_font("Helvetica", "I", 8)
    pdf.set_text_color(100, 100, 100)
    pdf.multi_cell(0, 4, pdf_safe("Avvertenza: " + DISCLAIMER), **nl)
    return bytes(pdf.output())


# ----------------------------------------------------------------------------
# WORD
# ----------------------------------------------------------------------------
def genera_docx(ctx, relazione, grafici=None):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    grafici = grafici if grafici is not None else grafici_report(ctx)
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(2)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)
    sec.footer.paragraphs[0].text = (
        f"ID report {ctx['report_id']} | {ctx['generato_il']} | Modello AI: {ctx['modello']} | FCR General Contractor / Aetheris Labs"
    )

    def tabella(intest, righe, destra_da=1):
        t = doc.add_table(rows=1, cols=len(intest))
        t.style = "Light Grid Accent 1"
        for i, h in enumerate(intest):
            t.rows[0].cells[i].text = h
        for r in righe:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = str(v)
                if i >= destra_da:
                    cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(8.5)
        doc.add_paragraph()

    def immagine(nome, larghezza=16):
        png = grafici.get(nome)
        if png:
            doc.add_picture(io.BytesIO(png), width=Cm(larghezza))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    base, scen, rating, piano = ctx["dati"], ctx["scen"], ctx["rating"], ctx["piano"]
    doc.add_heading("Valutazione del merito di credito e business plan", level=0)
    doc.add_paragraph("FCR General Contractor - Risk Management Report (analisi indicativa per il Comitato Fidi)")
    if ctx.get("pratica"):
        doc.add_paragraph(f"Pratica: {ctx['pratica']}")
    if ctx.get("dati_demo"):
        p = doc.add_paragraph()
        r = p.add_run("ATTENZIONE: il report e' basato su dati di esempio non verificati.")
        r.bold = True

    doc.add_heading(f"1. Sintesi indicatori (anno base {int(base['anno_riferimento'])})", level=1)
    tabella(
        ["Parametro", "Base", "Scenario"],
        [
            ("Ricavi delle vendite", fmt_eur(base["ricavi"]), fmt_eur(scen["ricavi"])),
            ("EBITDA (MOL)", fmt_eur(base["ebitda"]), fmt_eur(scen["ebitda"])),
            ("Patrimonio netto", fmt_eur(base["patrimonio_netto"]), fmt_eur(scen["pn"])),
            ("Posizione finanziaria netta", fmt_eur(base["pfn"]), fmt_eur(scen["pfn"])),
            ("PFN / EBITDA", fmt_x(ratio(base["pfn"], base["ebitda"])), fmt_x(rating["pfn_ebitda"])),
            ("PFN / Patrimonio netto", fmt_x(ratio(base["pfn"], base["patrimonio_netto"])), fmt_x(rating["pfn_pn"])),
            ("DSCR (anno 1 del piano)", "-", fmt_x(rating["dscr"]) if rating["dscr"] is not None else "n.d."),
            ("Punteggio (max 8)", "-", str(rating["punti"])),
        ],
    )
    p = doc.add_paragraph()
    p.add_run("Giudizio stimato (scenario): ").bold = True
    p.add_run(rating["classe"])
    tabella(
        ["Indicatore", "Valore", "Soglie (verde / giallo)", "Punti"],
        [(d["indicatore"], _valore_dettaglio(d), _soglie_testo(d), f"{d['punti']}/2") for d in rating["dettaglio"]],
    )

    doc.add_heading("2. Proiezioni pluriennali", level=1)
    tabella(
        ["Anno", "Ricavi", "EBITDA", "FCF", "PFN", "PFN/EBITDA", "DSCR"],
        [(r["Anno"], fmt_num(r["Ricavi"]), fmt_num(r["EBITDA"]), fmt_num(r["FCF"]), fmt_num(r["PFN"]),
          fmt_x(r["PFN/EBITDA"]), fmt_x(r["DSCR"]) if r["DSCR"] is not None else "n.d.") for r in piano],
    )
    p = ctx["piano_par"]
    doc.add_paragraph(
        f"Ipotesi: crescita ricavi {fmt_num(p['g_ricavi'], 1)}%, EBITDA {fmt_num(p['g_ebitda'], 1)}%, capex {fmt_num(p['capex'])} EUR, "
        f"ammortamenti {fmt_num(p['amm'])} EUR, rimborso annuo {fmt_num(p['rimborso'])} EUR, tasso {fmt_num(p['tasso'], 2)}%, "
        f"aliquota {fmt_num(p['aliquota'], 1)}%, CCN {fmt_num(p['ccn_pct'], 0)}% della variazione dei ricavi."
    )
    for nome in ("pfn_ebitda", "dscr", "cascata", "storico", "confronto"):
        immagine(nome)

    if ctx.get("cov_res"):
        doc.add_heading("Verifica covenant", level=1)
        tabella(
            ["Anno", "PFN/EBITDA", "Limite", "Esito", "DSCR", "Limite", "Esito"],
            [(c["Anno"], fmt_x(c["PFN/EBITDA"]), fmt_x(c["Limite PFN/EBITDA"]), c["Esito leva"],
              fmt_x(c["DSCR"]) if c["DSCR"] is not None else "n.d.", fmt_x(c["Limite DSCR"]), c["Esito DSCR"]) for c in ctx["cov_res"]],
        )
    if ctx.get("stress"):
        doc.add_heading("Stress test", level=1)
        tabella(
            ["Caso", "PFN/EBITDA anno 3", "DSCR minimo", "Punti", "Fascia"],
            [(s["Caso"], fmt_x(s["PFN/EBITDA anno 3"]), fmt_x(s["DSCR minimo"]) if s["DSCR minimo"] is not None else "n.d.",
              f"{s['Punteggio']}/8", s["Fascia"]) for s in ctx["stress"]],
        )
    k = ctx.get("commesse_kpi")
    if k:
        doc.add_heading("Portafoglio commesse", level=1)
        tabella(
            ["Indicatore", "Valore"],
            [
                ("Numero commesse", str(k["n_commesse"])),
                ("Importo totale contratti", fmt_eur(k["importo_totale"])),
                ("Backlog residuo", fmt_eur(k["backlog"])),
                ("Copertura ricavi (backlog / ricavi base)", fmt_x(k["copertura_ricavi"]) if k["copertura_ricavi"] is not None else "n.d."),
                ("Produzione eseguita non fatturata (negativo = anticipi)", fmt_eur(k["da_fatturare"])),
                ("Crediti da commesse", fmt_eur(k["crediti"])),
                ("Margine medio ponderato sul backlog", fmt_pct(k["margine_backlog_pct"])),
                ("Primo cliente / primi 3", f"{fmt_pct(k['top1_pct'], 0)} / {fmt_pct(k['top3_pct'], 0)}"),
            ],
        )
        immagine("commesse")

    doc.add_heading("3. Relazione e parere del Risk Management", level=1)
    for blocco in (relazione or "").split("\n\n"):
        if blocco.strip():
            doc.add_paragraph(blocco.strip())
    p = doc.add_paragraph()
    r = p.add_run("Avvertenza: " + DISCLAIMER)
    r.italic = True
    r.font.size = Pt(8)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------------------
# EXCEL CON FORMULE VIVE
# ----------------------------------------------------------------------------
def genera_excel(ctx):
    """
    Cartella di lavoro con formule: il foglio 'Input' contiene le celle modificabili (gialle),
    'Piano' e 'Rating' si ricalcolano in Excel. Riproduce la stessa logica di core.py.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    giallo = PatternFill("solid", fgColor="FFF2CC")
    testata = PatternFill("solid", fgColor="1E293B")
    f_in = Font(color="1D4ED8", bold=True)
    f_h = Font(color="FFFFFF", bold=True)
    f_tit = Font(bold=True, size=13)
    EUR = '#,##0;[Red]-#,##0'
    XF = '0.00"x"'

    def input_cell(ws, ref, value, fmt=None):
        ws[ref] = value
        ws[ref].fill = giallo
        ws[ref].font = f_in
        if fmt:
            ws[ref].number_format = fmt

    def sezione(ws, riga, testo, colonne=3):
        for c in range(1, colonne + 1):
            ws.cell(row=riga, column=c).fill = testata
        ws.cell(row=riga, column=1, value=testo).font = f_h

    d, leve, p = ctx["dati"], ctx["leve"], ctx["piano_par"]
    s = ctx.get("soglie") or SOGLIE_DEFAULT
    cov = ctx.get("covenant") or {"pfn_ebitda_max": 3.5, "dscr_min": 1.2}

    # --- Input --------------------------------------------------------------
    ws = wb.active
    ws.title = "Input"
    ws["A1"] = "FCR General Contractor - Input del modello (celle gialle modificabili)"
    ws["A1"].font = f_tit
    sezione(ws, 2, "DATI DI BILANCIO (anno base)")
    voci = [
        (3, "Anno di riferimento", d["anno_riferimento"], "0"),
        (4, "Ricavi (EUR)", d["ricavi"], EUR),
        (5, "EBITDA (EUR)", d["ebitda"], EUR),
        (6, "Utile netto (EUR)", d["utile_netto"], EUR),
        (7, "Totale attivo (EUR)", d["totale_attivo"], EUR),
        (8, "Patrimonio netto (EUR)", d["patrimonio_netto"], EUR),
        (9, "Debiti finanziari totali (EUR)", d["debiti_finanziari_totali"], EUR),
        (10, "Posizione finanziaria netta - PFN (EUR)", d["pfn"], EUR),
    ]
    for r, lab, val, fmt in voci:
        ws.cell(row=r, column=1, value=lab)
        input_cell(ws, f"B{r}", val, fmt)
    sezione(ws, 12, "LEVE DI SCENARIO")
    for r, lab, val in [
        (13, "Variazione debito netto - PFN (EUR, + = piu' debito)", leve["delta_pfn"]),
        (14, "Iniezione di capitale (EUR, aumenta PN e riduce PFN)", leve["delta_pn"]),
        (15, "Variazione EBITDA (EUR)", leve["delta_ebitda"]),
    ]:
        ws.cell(row=r, column=1, value=lab)
        input_cell(ws, f"B{r}", val, EUR)
    sezione(ws, 17, "PARAMETRI DEL PIANO")
    for r, lab, val, fmt in [
        (18, "Crescita annua ricavi (%)", p["g_ricavi"], "0.0"),
        (19, "Crescita annua EBITDA (%)", p["g_ebitda"], "0.0"),
        (20, "CCN su variazione ricavi (%)", p["ccn_pct"], "0.0"),
        (21, "Capex annuo (EUR)", p["capex"], EUR),
        (22, "Ammortamenti annui (EUR)", p["amm"], EUR),
        (23, "Rimborso annuo quota capitale (EUR)", p["rimborso"], EUR),
        (24, "Tasso d'interesse sulla PFN (%)", p["tasso"], "0.00"),
        (25, "Aliquota fiscale (%)", p["aliquota"], "0.0"),
    ]:
        ws.cell(row=r, column=1, value=lab)
        input_cell(ws, f"B{r}", val, fmt)
    sezione(ws, 27, "SOGLIE DI SCORING (colonna B = verde, colonna C = giallo)")
    for r, lab, k, fmt in [
        (28, "PFN / EBITDA (minore = meglio)", "pfn_ebitda", "0.00"),
        (29, "PFN / Patrimonio netto (minore = meglio)", "pfn_pn", "0.00"),
        (30, "Patrimonio netto / Attivo (decimale, maggiore = meglio)", "equity_ratio", "0.00"),
        (31, "DSCR (maggiore = meglio)", "dscr", "0.00"),
    ]:
        ws.cell(row=r, column=1, value=lab)
        input_cell(ws, f"B{r}", s[k][0], fmt)
        input_cell(ws, f"C{r}", s[k][1], fmt)
    sezione(ws, 33, "COVENANT")
    ws["A34"] = "PFN / EBITDA massimo"
    input_cell(ws, "B34", cov["pfn_ebitda_max"], "0.00")
    ws["A35"] = "DSCR minimo"
    input_cell(ws, "B35", cov["dscr_min"], "0.00")
    ws.column_dimensions["A"].width = 58
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 14

    # --- Piano --------------------------------------------------------------
    wp = wb.create_sheet("Piano")
    wp["A1"] = "Piano pluriennale (formule collegate al foglio Input)"
    wp["A1"].font = f_tit
    wp["A3"], wp["B3"] = "EBITDA scenario", "=Input!B5+Input!B15"
    wp["A4"], wp["B4"] = "PFN scenario", "=Input!B10+Input!B13-Input!B14"
    wp["A5"], wp["B5"] = "Patrimonio netto scenario", "=Input!B8+Input!B14"
    wp["A6"], wp["B6"] = "Attivo scenario", "=Input!B7+Input!B14"
    wp["A7"], wp["B7"] = "Ricavi", "=Input!B4"
    for r in range(3, 8):
        wp[f"B{r}"].number_format = EUR
    for c in range(1, 6):
        wp.cell(row=9, column=c).fill = testata
        wp.cell(row=9, column=c).font = f_h
    wp["A9"] = "Voce (EUR)"
    wp["B9"] = "Anno base"
    for i, col in enumerate("CDE", start=1):
        wp[f"{col}9"] = f'="Anno {i} ("&(Input!$B$3+{i})&")"'
    etichette = {
        10: "Ricavi", 11: "EBITDA", 12: "Interessi", 13: "Imposte", 14: "Variazione CCN", 15: "Capex",
        16: "Free cash flow", 17: "PFN", 18: "PFN / EBITDA", 19: "Servizio del debito", 20: "DSCR",
        22: "Covenant PFN / EBITDA", 23: "Covenant DSCR",
    }
    for r, lab in etichette.items():
        wp.cell(row=r, column=1, value=lab)
    wp["B10"], wp["B11"], wp["B17"] = "=B7", "=B3", "=B4"
    cols = ["B", "C", "D", "E"]
    for i in range(1, 4):
        c, pr = cols[i], cols[i - 1]
        wp[f"{c}10"] = f"={pr}10*(1+Input!$B$18/100)"
        wp[f"{c}11"] = f"={pr}11*(1+Input!$B$19/100)"
        wp[f"{c}12"] = f"=MAX({pr}17,0)*Input!$B$24/100"
        wp[f"{c}13"] = f"=MAX({c}11-Input!$B$22-{c}12,0)*Input!$B$25/100"
        wp[f"{c}14"] = f"=Input!$B$20/100*({c}10-{pr}10)"
        wp[f"{c}15"] = "=Input!$B$21"
        wp[f"{c}16"] = f"={c}11-{c}13-{c}15-{c}12-{c}14"
        wp[f"{c}17"] = f"={pr}17-{c}16"
        wp[f"{c}19"] = f"={c}12+Input!$B$23"
        wp[f"{c}20"] = f'=IF({c}19>0,({c}11-{c}13-{c}14)/{c}19,"n.d.")'
        wp[f"{c}22"] = f'=IF(ISNUMBER({c}18),IF({c}18<=Input!$B$34,"OK","SFORATO"),IF({c}18="n.s.","SFORATO","OK"))'
        wp[f"{c}23"] = f'=IF(ISNUMBER({c}20),IF({c}20>=Input!$B$35,"OK","SFORATO"),"OK")'
    for c in cols:
        wp[f"{c}18"] = f'=IF({c}11>0,{c}17/{c}11,IF({c}17>0,"n.s.",0))'
        for r in (10, 11, 12, 13, 14, 15, 16, 17, 19):
            wp[f"{c}{r}"].number_format = EUR
        wp[f"{c}18"].number_format = XF
        wp[f"{c}20"].number_format = XF
        for r in (18, 20, 22, 23):
            wp[f"{c}{r}"].alignment = Alignment(horizontal="right")
    wp["A25"] = ("FCF = EBITDA - imposte - capex - interessi - variazione CCN; PFN(t) = PFN(t-1) - FCF. "
                 "DSCR = (EBITDA - imposte - variazione CCN) / (interessi + rimborso).")
    wp.column_dimensions["A"].width = 30
    for c in "BCDE":
        wp.column_dimensions[c].width = 18

    # --- Rating -------------------------------------------------------------
    wr = wb.create_sheet("Rating")
    wr["A1"] = "Scoring indicativo (formule)"
    wr["A1"].font = f_tit
    for c, h in enumerate(["Indicatore", "Valore", "Soglia verde", "Soglia giallo", "Punti"], start=1):
        wr.cell(row=3, column=c, value=h).fill = testata
        wr.cell(row=3, column=c).font = f_h
    wr["A4"], wr["A5"], wr["A6"], wr["A7"] = "PFN / EBITDA", "PFN / Patrimonio netto", "Patrimonio netto / Attivo", "DSCR (anno 1)"
    wr["B4"] = '=IF(Piano!B3>0,Piano!B4/Piano!B3,IF(Piano!B4>0,"n.s.",0))'
    wr["B5"] = '=IF(Piano!B5>0,Piano!B4/Piano!B5,IF(Piano!B4>0,"n.s.",0))'
    wr["B6"] = "=IF(Piano!B6>0,Piano!B5/Piano!B6,0)"
    wr["B7"] = "=Piano!C20"
    for r, src in zip((4, 5, 6, 7), (28, 29, 30, 31)):
        wr[f"C{r}"] = f"=Input!B{src}"
        wr[f"D{r}"] = f"=Input!C{src}"
    wr["E4"] = '=IF(ISNUMBER(B4),IF(B4<C4,2,IF(B4<D4,1,0)),0)'
    wr["E5"] = '=IF(ISNUMBER(B5),IF(B5<C5,2,IF(B5<D5,1,0)),0)'
    wr["E6"] = "=IF(B6>=C6,2,IF(B6>=D6,1,0))"
    wr["E7"] = '=IF(ISNUMBER(B7),IF(B7>=C7,2,IF(B7>=D7,1,0)),2)'
    for r in (4, 5, 7):
        wr[f"B{r}"].number_format = XF
    wr["B6"].number_format = "0.0%"
    for r in range(4, 8):
        wr[f"B{r}"].alignment = Alignment(horizontal="right")
    wr["A9"], wr["B9"] = "Punteggio totale (max 8)", "=SUM(E4:E7)"
    wr["A10"] = "Fascia"
    wr["B10"] = '=IF(OR(Piano!B5<=0,Piano!B3<=0,B9<4),"FASCIA ROSSA",IF(B9>=7,"FASCIA VERDE","FASCIA GIALLA"))'
    wr["A9"].font = wr["A10"].font = Font(bold=True)
    wr["B9"].font = wr["B10"].font = Font(bold=True)
    wr["A12"] = "Verde: punteggio >= 7. Rossa: punteggio < 4, oppure PN <= 0, oppure EBITDA <= 0. Modello indicativo, non regolamentare."
    wr.column_dimensions["A"].width = 32
    for c in "BCDE":
        wr.column_dimensions[c].width = 16

    # --- Storico (valori) ---------------------------------------------------
    wh = wb.create_sheet("Storico")
    wh["A1"] = "Storico esercizi (valori inseriti, non formule)"
    wh["A1"].font = f_tit
    for c, h in enumerate(["Anno", "Ricavi", "EBITDA", "Patrimonio netto", "PFN"], start=1):
        wh.cell(row=3, column=c, value=h).fill = testata
        wh.cell(row=3, column=c).font = f_h
    righe_st = list(ctx.get("storico") or []) + [
        {"anno": int(d["anno_riferimento"]), "ricavi": d["ricavi"], "ebitda": d["ebitda"],
         "patrimonio_netto": d["patrimonio_netto"], "pfn": d["pfn"]}
    ]
    for i, r in enumerate(righe_st, start=4):
        wh.cell(row=i, column=1, value=r["anno"])
        for j, k in enumerate(["ricavi", "ebitda", "patrimonio_netto", "pfn"], start=2):
            wh.cell(row=i, column=j, value=r[k]).number_format = EUR
    for c in "ABCDE":
        wh.column_dimensions[c].width = 18

    # --- Commesse (con formule) ----------------------------------------------
    cdf = ctx.get("commesse_df")
    if cdf is not None and len(cdf) > 0:
        wc = wb.create_sheet("Commesse")
        wc["A1"] = "Portafoglio commesse (colonne H-K e totali in formula)"
        wc["A1"].font = f_tit
        intest = ["Commessa", "Cliente", "Importo", "Avanzamento %", "Fatturato", "Incassato", "Margine %",
                  "Backlog", "Produzione eseguita", "Da fatturare", "Crediti"]
        for c, h in enumerate(intest, start=1):
            wc.cell(row=3, column=c, value=h).fill = testata
            wc.cell(row=3, column=c).font = f_h
        n = 0
        for i, (_, r) in enumerate(cdf.iterrows(), start=4):
            n += 1
            wc.cell(row=i, column=1, value=str(r["Commessa"]))
            wc.cell(row=i, column=2, value=str(r["Cliente"]))
            for j, k in enumerate(["Importo", "Avanzamento_pct", "Fatturato", "Incassato", "Margine_pct"], start=3):
                cell = wc.cell(row=i, column=j, value=float(r[k]) if r[k] == r[k] else 0.0)
                cell.fill, cell.font = giallo, f_in
                cell.number_format = "0.0" if k in ("Avanzamento_pct", "Margine_pct") else EUR
            wc[f"H{i}"] = f"=C{i}*(1-D{i}/100)"
            wc[f"I{i}"] = f"=C{i}*D{i}/100"
            wc[f"J{i}"] = f"=I{i}-E{i}"
            wc[f"K{i}"] = f"=E{i}-F{i}"
            for col in "HIJK":
                wc[f"{col}{i}"].number_format = EUR
        ult = 3 + n
        tot = ult + 2
        wc[f"A{tot}"] = "TOTALI"
        wc[f"A{tot}"].font = Font(bold=True)
        for col in "CEFHIJK":
            wc[f"{col}{tot}"] = f"=SUM({col}4:{col}{ult})"
            wc[f"{col}{tot}"].number_format = EUR
            wc[f"{col}{tot}"].font = Font(bold=True)
        wc[f"A{tot + 1}"] = "Margine medio ponderato sul backlog (%)"
        wc[f"H{tot + 1}"] = f"=IF(H{tot}>0,SUMPRODUCT(H4:H{ult},G4:G{ult})/H{tot},0)"
        wc[f"H{tot + 1}"].number_format = "0.0"
        wc[f"A{tot + 2}"] = "Copertura ricavi (backlog / ricavi base)"
        wc[f"H{tot + 2}"] = f"=IF(Input!B4>0,H{tot}/Input!B4,0)"
        wc[f"H{tot + 2}"].number_format = XF
        for c in range(1, 12):
            wc.column_dimensions[get_column_letter(c)].width = 17

    # --- Note ---------------------------------------------------------------
    wn = wb.create_sheet("Note")
    righe_note = [
        "FCR General Contractor / Aetheris Labs - cartella di lavoro generata dall'applicazione",
        f"ID report: {ctx['report_id']}  |  Generato il: {ctx['generato_il']}  |  Modello AI: {ctx['modello']}",
        "Le celle gialle (foglio Input e Commesse) sono modificabili; Piano e Rating si ricalcolano.",
        "n.s. = non significativo (EBITDA o patrimonio netto non positivi con debito positivo). n.d. = non disponibile.",
        "I valori dello storico sono inseriti come numeri (non formule).",
        "",
        "AVVERTENZA: " + DISCLAIMER,
    ]
    if ctx.get("dati_demo"):
        righe_note.insert(2, "ATTENZIONE: dati di esempio non verificati.")
    for i, t in enumerate(righe_note, start=1):
        wn.cell(row=i, column=1, value=t)
        wn.cell(row=i, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    wn.column_dimensions["A"].width = 120

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------------------
# COSTRUZIONE DEL CONTESTO (usata da app e test)
# ----------------------------------------------------------------------------
def costruisci_ctx(dati, leve, piano_par, soglie=None, covenant=None, storico=None, commesse_df=None,
                   scenari_salvati=None, dati_demo=False, modello="n.d.", pratica=""):
    """Calcola tutto cio' che serve ai report a partire dagli input."""
    import time
    import core

    soglie = soglie or core.SOGLIE_DEFAULT
    covenant = covenant or core.COVENANT_DEFAULT
    res = core.esegui_scenario(dati, leve, piano_par, soglie)
    kpi, per_cliente = core.analisi_commesse(commesse_df, dati["ricavi"])

    scenari = {"Base": core.esegui_scenario(dati, core.LEVE_DEFAULT, piano_par, soglie)}
    for nome, sc in (scenari_salvati or {}).items():
        scenari[nome] = core.esegui_scenario(dati, sc.get("leve", {}), sc.get("piano_par", piano_par), soglie)
    if len(scenari) == 1:
        scenari = {}

    firma = {
        "dati": dati, "leve": leve, "piano": piano_par, "soglie": soglie, "cov": covenant,
        "storico": storico or [], "scenari": scenari_salvati or {},
        "commesse": commesse_df.to_dict("records") if commesse_df is not None else [],
    }
    return {
        "dati": dati, "scen": res["scen"], "piano": res["piano"], "rating": res["rating"],
        "leve": res["leve"], "piano_par": piano_par, "soglie": soglie, "covenant": covenant,
        "cov_res": core.verifica_covenant(res["piano"], covenant),
        "stress": core.stress_test(dati, leve, piano_par, soglie),
        "commesse_kpi": kpi, "commesse_df": commesse_df if kpi else None,
        "per_cliente": per_cliente, "storico": storico or [], "scenari": scenari,
        "dati_demo": dati_demo, "report_id": core.hash_payload(firma, 10).upper(),
        "hash_dati": core.hash_payload(dati), "modello": modello,
        "generato_il": time.strftime("%d/%m/%Y %H:%M"), "pratica": pratica,
    }
