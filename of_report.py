"""
of_report.py - Offerta commerciale FCR in PPTX e PDF dallo stesso schema di diapositive.

`costruisci(ctx)` produce l'elenco delle diapositive (dati + posizione), `pptx_bytes` e `pdf_bytes` le disegnano.
Stile: titolo blu scuro, schede bianche, Opzione A verde, Opzione B oro, accenti azzurri (come l'offerta attuale).
"""
import io
import os

import of

QUI = os.path.dirname(os.path.abspath(__file__))
W, H = 13.333, 7.5
COL = {"titolo": "163247", "sotto": "647686", "testo": "253544", "verde": "5A8F72", "oro": "D4A84A", "navy": "0E2536", "azzurro": "1A8ACB", "bianco": "FFFFFF",
       "sfondo": "F8FAFC", "bordo": "DDE3EA", "verde_c": "F0F8F2", "oro_c": "FFF8EA", "azz_c": "EAF3FB", "rosa_c": "FFF4F1", "grigio_c": "F1F5F9", "rosso": "A94B4B"}
LOGO = next((p for p in (os.path.join(QUI, "fcr_logo.png"),) if os.path.exists(p)), None)
FONT_R, FONT_B = os.path.join(QUI, "fonts", "LiberationSans-Regular.ttf"), os.path.join(QUI, "fonts", "LiberationSans-Bold.ttf")
NOTA_FIN = "La simulazione è meramente indicativa e non costituisce delibera di finanziamento. Tasso, spese, condizioni, istruttoria e merito creditizio restano soggetti alla valutazione dell'ente finanziatore."
NOTA_FISC = "Le percentuali di detrazione sono scenari di simulazione. Effettiva spettanza, limiti, capienza fiscale e requisiti devono essere verificati per ciascun contribuente."


def _it(v, d=2):
    return f"{v:,.{d}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _e(v, d=2):
    return f"{_it(v, d)} €"


# ----------------------------------------------------------------------------------------------------------------------------- elementi
def kpi(label, valore, nota="", colore="azzurro"):
    return {"k": "kpi", "label": label, "valore": valore, "nota": nota, "colore": colore}


def card(titolo, testo, fondo="bianco", colore="titolo"):
    return {"k": "card", "titolo": titolo, "testo": testo, "fondo": fondo, "colore": colore}


def tabella(colonne, righe, pesi=None, colore="navy", evidenzia_ultima=False):
    return {"k": "tab", "colonne": colonne, "righe": righe, "pesi": pesi or [1] * len(colonne), "colore": colore, "ultima": evidenzia_ultima}


def nota(testo, fondo="azz_c"):
    return {"k": "nota", "testo": testo, "fondo": fondo}


def passo(n, titolo, testo, colore="azzurro"):
    return {"k": "passo", "n": str(n), "titolo": titolo, "testo": testo, "colore": colore}


def riga(h, *els, pesi=None):
    return {"h": h, "els": list(els), "pesi": pesi or [1] * len(els)}


def slide(titolo, sotto, *righe, fonte=""):
    return {"tipo": "std", "titolo": titolo, "sotto": sotto, "righe": list(righe), "fonte": fonte}


# ----------------------------------------------------------------------------------------------------------------------------- contenuto
def costruisci(ctx):
    p = ctx.get("pratica", {})
    cond = p.get("condominio") or "il Condominio"
    opz = [o for o in ("A", "B") if ctx.get("opzioni", {}).get(o, {}).get("imponibile", 0) > 0]
    colore = {"A": "verde", "B": "oro"}
    S = []
    unita = ctx.get("unita") or []
    n_ui = int(p.get("ui") or len(unita) or 0)
    validita = int(p.get("validita_gg") or 15)
    sotto_cop = " | ".join(x for x in (f"{len(opz)} soluzioni tecniche" if len(opz) > 1 else "", f"{n_ui} unità abitative" if n_ui else "",
                                        f"finanziamento da {min(ctx['params']['durate'])} a {max(ctx['params']['durate'])} anni" if ctx.get("params") else "") if x)
    S.append({"tipo": "copertina", "titolo": f"Proposta di riqualificazione\n{cond}", "riga1": " | ".join(x for x in (p.get("indirizzo"), p.get("cod")) if x),
              "riga2": sotto_cop, "data": f"{p.get('data', '')}\nVALIDITÀ OFFERTA: {validita} GIORNI"})
    S.append(slide("Riservatezza e limiti di utilizzo", "Avvertenza parte integrante della proposta e degli allegati.",
                   riga(1.5, card("Destinazione della proposta", f"La presente proposta, i prezzi, le soluzioni tecniche, i computi, le simulazioni finanziarie e gli elaborati FCR sono messi a disposizione esclusivamente di {cond}, del suo amministratore e dei condomini in sede assembleare."),
                        card("Divieto di diffusione competitiva", "La documentazione non è destinata alla trasmissione a imprese concorrenti, fornitori o soggetti terzi per richieste di offerte comparative o al ribasso. Qualsiasi ulteriore utilizzo deve essere previamente autorizzato per iscritto da FCR.")),
                   riga(1.5, card("Tutela del contenuto", "Restano riservati il know-how commerciale, le modalità di strutturazione dell'offerta, le analisi economiche, le ipotesi tecniche e i documenti elaborati da FCR."),
                        card("Condivisione consentita", "È consentita la circolazione interna all'assemblea condominiale e ai professionisti incaricati dal Condominio per la sola valutazione tecnica, amministrativa e finanziaria dell'intervento.")),
                   riga(0.6, nota("DOCUMENTO RISERVATO - USO ESCLUSIVO DEL DESTINATARIO E DELL'ASSEMBLEA CONDOMINIALE", "rosa_c"))))
    S.append(slide("FCR General Contractor", "Un unico interlocutore per progettazione, cantiere, controllo e gestione economica.",
                   riga(1.1, kpi("PATRIMONIO NETTO", "> € 15 mln", "", "azzurro"), kpi("OPERATIVITÀ", "General Contractor", "", "verde"), kpi("MODELLO", "Progetto + cantiere", "", "oro")),
                   riga(1.5, card("Coordinamento unitario", "FCR integra la lettura tecnica del progetto, la cantierizzazione, le lavorazioni specialistiche, la gestione della sicurezza e il controllo economico in un'unica catena di responsabilità operativa."),
                        card("Controllo dei costi", "La contabilità a misura consente di ricondurre il corrispettivo alle quantità effettivamente eseguite, mantenendo tracciabilità su lavorazioni, varianti e imprevisti."),
                        card("Capacità di esecuzione", "Le lavorazioni specialistiche vengono coordinate secondo una sequenza unica di cantiere, con interfaccia centralizzata verso il Condominio.")),
                   riga(0.9, nota("Struttura della proposta: quadro economico completo, simulazione di finanziamento a rata costante, lettura del beneficio fiscale, energetico e patrimoniale."))))
    S.append(slide("Come FCR governa l'intervento", "Sequenza operativa costruita per ridurre frammentazione, ritardi e contenziosi tecnici.",
                   riga(0.72, passo(1, "Progetto e verifica", "Lettura del computo, controlli quantità, dettagli costruttivi e interferenze.")),
                   riga(0.72, passo(2, "Cantierizzazione", "Ponteggi, sicurezza, accessi, protezioni e logistica.")),
                   riga(0.72, passo(3, "Esecuzione", "Demolizioni, impermeabilizzazioni, coibentazioni, coperture, lattonerie.")),
                   riga(0.72, passo(4, "Finiture", "Facciate, parapetti, ripristini murari e opere accessorie.")),
                   riga(0.72, passo(5, "Contabilità", "Misure effettive, SAL, autorizzazioni e chiusura economica.")),
                   riga(0.6, nota("Obiettivo: trasformare un intervento complesso in un processo unico, misurabile e deliberabile."))))
    # sintesi
    k = []
    for o in opz:
        d = ctx["opzioni"][o]
        k.append(kpi(f"OPZIONE {o} - TOTALE PROGETTO", _e(d["qte"]["totale"]), d.get("nome", ""), colore[o]))
    k_riga = [kpi("UNITÀ ABITATIVE", str(n_ui), f"media teorica {_it(1000 / n_ui)} millesimi" if n_ui else "", "azzurro")] + k
    k_riga2 = [kpi("VALIDITÀ", f"{validita} giorni", "dalla data di emissione", "azzurro")]
    if ctx.get("params"):
        k_riga2.append(kpi("FINANZIAMENTO", " / ".join(f"{a} anni" for a in ctx["params"]["durate"]), f"rata costante, TAN {_it(ctx['params']['tan'] * 100)}%", "navy"))
    S.append(slide("Sintesi della proposta", p.get("sintesi_sotto") or "Perimetro di opere e alternative tecniche a confronto.",
                   riga(1.15, *k_riga), riga(1.15, *k_riga2),
                   riga(2.2, *[card(f"Opzione {o}", ctx["opzioni"][o].get("desc") or "", "verde_c" if o == "A" else "oro_c", colore[o]) for o in opz]),
                   riga(0.9, card("Perimetro comune", p.get("opere_comuni") or "Cantierizzazione, ponteggi, demolizioni e preparazioni, lattonerie, tinteggiature e ripristini murari a misura.", "grigio_c"))))
    if p.get("stato_fatto"):
        S.append(slide("Stato di fatto e obiettivi", "Perché intervenire e cosa si ottiene.", riga(2.4, card("Stato di fatto", p["stato_fatto"], "rosa_c", "rosso"), card("Obiettivi dell'intervento", p.get("descrizione") or "", "azz_c", "azzurro"))))
    elif p.get("descrizione"):
        S.append(slide("L'intervento proposto", "Descrizione delle lavorazioni.", riga(3.0, card("Descrizione", p["descrizione"], "azz_c", "azzurro"))))
    if p.get("opere_comuni_dettaglio"):
        S.append(slide("Opere comuni alle alternative", "Il perimetro di preparazione e finitura è lo stesso.", riga(3.5, card("Opere comuni", p["opere_comuni_dettaglio"], "grigio_c"))))
    for o in opz:
        d = ctx["opzioni"][o]
        S.append(slide(f"Opzione {o} - {d.get('nome') or 'soluzione tecnica'}", "Descrizione tecnica e importo.",
                       riga(2.3, card("Descrizione", d.get("desc") or "", "verde_c" if o == "A" else "oro_c", colore[o])),
                       riga(1.15, kpi("IMPORTO LAVORI IVA ESCLUSA", _e(d["imponibile"]), "", colore[o]), kpi("TOTALE PROGETTO con IVA e spese tecniche", _e(d["qte"]["totale"]), "", "navy"))))
    if len(opz) > 1 and ctx.get("confronto"):
        dif = ctx["opzioni"]["A"]["imponibile"] - ctx["opzioni"]["B"]["imponibile"]
        righe_c = [[r.get("Parametro", ""), r.get("A", ""), r.get("B", "")] for r in ctx["confronto"] if r.get("Parametro")]
        righe_c.append(["Costo lavori IVA esclusa", _e(ctx["opzioni"]["A"]["imponibile"]), _e(ctx["opzioni"]["B"]["imponibile"])])
        S.append(slide("Confronto tecnico tra le due opzioni", "Cambia il modo in cui viene costruito il nuovo involucro.",
                       riga(0.4 * (len(righe_c) + 1) + 0.1, tabella(["Parametro", "Opzione A", "Opzione B"], righe_c, [1.4, 1, 1], evidenzia_ultima=True)),
                       riga(0.6, nota(f"Differenza di costo lavori: {_e(abs(dif))} ({'A più cara di B' if dif > 0 else 'B più cara di A'}, {_it(abs(dif) / min(ctx['opzioni']['A']['imponibile'], ctx['opzioni']['B']['imponibile']) * 100)}%). La scelta deve essere tecnica prima che economica."))))
    n_mill = of.totale_millesimi(unita)
    for o in opz:
        q = ctx["opzioni"][o]["qte"]
        S.append(slide(f"Quadro economico - Opzione {o}", "Importi derivati dal computo metrico e dal quadro economico.",
                       riga(2.9, tabella(["Voce", "Importo"], [["Lavori imponibili", _e(q["imponibile"])], [f"IVA lavori {int(round(ctx['qte_par']['iva_lavori'] * 100))}%", _e(q["iva_lavori"])],
                                                              ["Totale lavori con IVA", _e(q["lavori_iva"])], ["Spese professionali + Responsabile Lavori", _e(q["spese"])], ["TOTALE PROGETTO", _e(q["totale"])]],
                                         [3, 1], colore[o], True),
                            kpi("CAPITALE PER MILLESIMO", _e(q["totale"] / 1000), "prima degli interessi del finanziamento", colore[o]),
                            pesi=[2.2, 1]),
                       riga(1.0, kpi(f"MEDIA ARITMETICA PER {n_ui} UNITÀ" if n_ui else "MEDIA PER UNITÀ", _e(q["totale"] / n_ui) if n_ui else "-", "solo indicativa: prevalgono i millesimi reali", "azzurro")),
                       riga(0.9, nota("La quota realmente imputabile a ciascuna unità dipende dalla tabella millesimale e dalle regole di riparto deliberate. La media serve soltanto a dare un ordine di grandezza assembleare."))))
    # condizioni
    S.append(slide("Condizioni economiche e natura dell'offerta", "", riga(1.7, card("Lavorazioni impreviste", p.get("imprevisti") or "Gli importi potranno variare per lavorazioni non prevedibili che si rendessero necessarie a seguito delle effettive condizioni dell'immobile, di elementi non visibili prima delle demolizioni o di prescrizioni tecniche e normative applicabili.", "rosa_c", "rosso")),
                   riga(1.6, card("A carico del Condominio", p.get("carico_condominio") or "Oneri di occupazione di suolo pubblico; fornitura di acqua ed energia elettrica di cantiere, secondo quanto indicato nel computo."),
                        card("A carico dei singoli condòmini", p.get("carico_condomini") or "Interventi interferenti sulle parti private, come previsto dal computo."),
                        card("Validità", f"{validita} giorni solari continuativi dalla data del preventivo. Decorso il termine, prezzi e condizioni potranno essere aggiornati.", "azz_c", "azzurro"))))
    # finanziamento
    par = ctx.get("params")
    sim = ctx.get("sim") or {}
    if par and sim:
        for o in opz:
            r = sim[o]
            righe_f = [[f"{x['anni']} anni ({x['mesi']} mesi)", _e(x["rata"]), _e(x["rata"] / 1000), _e(x["interessi"]), _e(x["totale"])] for x in r]
            S.append(slide(f"Finanziamento condominiale - Opzione {o}", f"Ammortamento alla francese, rata costante, TAN ipotizzato {_it(par['tan'] * 100)}% annuo. Nessun istituto finanziatore indicato.",
                           riga(0.4 * (len(righe_f) + 1) + 0.1, tabella(["Durata", "Rata condominio / mese", "Rata / millesimo / mese", "Interessi totali", "Totale rimborsato"], righe_f, [1.2, 1.3, 1.3, 1.2, 1.2], colore[o])),
                           riga(1.1, kpi("CAPITALE FINANZIATO", _e(ctx["opzioni"][o]["qte"]["totale"]), "totale progetto con IVA e spese tecniche", colore[o]),
                                kpi("MEDIA PER UNITÀ (rata più lunga)", _e(r[-1]["rata"] / n_ui) if n_ui else "-", "al mese, solo indicativa", "azzurro")),
                           riga(0.8, nota(NOTA_FIN, "grigio_c"))))
        # recupero fiscale
        righe_r = []
        for o in opz:
            q = ctx["opzioni"][o]["qte"]
            base = q["lavori_iva"] + (q["spese"] if par.get("base_con_spese") else 0)
            righe_r.append([f"Opzione {o}", _e(base)] + [_e(base * d / 100) for d in par["detrazioni"]])
        S.append(slide(f"Recupero fiscale complessivo su {par['anni']} anni", f"Beneficio fiscale su {'lavori con IVA e spese tecniche' if par.get('base_con_spese') else 'lavori comprensivi di IVA'}, nelle ipotesi di recupero al " + " oppure al ".join(f"{int(d)}%" for d in par["detrazioni"]) + ".",
                       riga(0.4 * (len(righe_r) + 1) + 0.1, tabella(["", "Base di calcolo ipotizzata"] + [f"Recupero {int(d)}%" for d in par["detrazioni"]], righe_r, [1, 1.4] + [1] * len(par["detrazioni"]))),
                       riga(1.1, kpi("ORIZZONTE TEMPORALE", f"{par['anni']} anni", "quote fiscali distribuite nel periodo", "azzurro"),
                            kpi("RISPARMIO ENERGETICO STIMATO", _e(par.get("risparmio_annuo", 0)) + " / anno" if par.get("risparmio_annuo") else "da valutare",
                                "dato dichiarato dal progettista / APE" if par.get("risparmio_annuo") else "non incluso nelle simulazioni", "verde")),
                       riga(0.7, nota(NOTA_FISC, "grigio_c"))))
        # costo reale / impatto
        for o in opz:
            r = sim[o]
            righe_i, ris = [], par.get("risparmio_annuo", 0)
            for x in r:
                for d in par["detrazioni"]:
                    dd = int(d)
                    righe_i.append([f"{x['anni']} anni", f"{dd}%", _e(x["rata"] / 1000), _e((x["rata"] - x[f"impatto{dd}"] - ris / 12) / 1000),
                                    _e(ris / 12 / 1000), _e(x[f"impatto{dd}"] / 1000), _e(x[f"costo_reale{dd}"] / 1000)])
            S.append(slide(f"Impatto economico reale per millesimo - Opzione {o}", "Impatto mensile = rata - recupero fiscale - risparmio energetico. Costo reale = capitale + interessi - recuperi.",
                           riga(0.38 * (len(righe_i) + 1) + 0.1, tabella(["Durata", "Detraz.", "Rata/mill./mese", "Recupero fiscale/mill./mese", "Risparmio energ./mill./mese", "IMPATTO REALE/mill./mese", "Costo reale/millesimo"], righe_i, [0.8, 0.7, 1, 1.3, 1.3, 1.3, 1.2], colore[o], False)),
                           riga(0.7, nota("La rata bancaria è l'uscita di cassa; il recupero fiscale e il risparmio energetico sono benefici distribuiti nel tempo: l'impatto reale li sottrae mese per mese."), )))
        # ripartizione per unita'
        sc = ctx.get("scenario") or {}
        if unita and sc.get("opzione") in opz:
            x = next((z for z in sim[sc["opzione"]] if z["anni"] == sc["anni"]), sim[sc["opzione"]][-1])
            dd = int(sc["detrazione"])
            ris_m = par.get("risparmio_annuo", 0) / 12
            rec_m = x["rata"] - x[f"impatto{dd}"] - ris_m
            tab = [[u["Unita"], u["Proprietario"], _it(u["Millesimi"], 3), _e(u["Millesimi"] / 1000 * ctx["opzioni"][sc["opzione"]]["qte"]["totale"]), _e(x["rata"] * u["Millesimi"] / 1000),
                    _e(rec_m * u["Millesimi"] / 1000), _e(ris_m * u["Millesimi"] / 1000), _e(x["impatto" + str(dd)] * u["Millesimi"] / 1000)] for u in unita]
            for i in range(0, len(tab), 11):
                chunk = tab[i:i + 11]
                S.append(slide(f"Ripartizione per unità - Opzione {sc['opzione']}, {sc['anni']} anni, detrazione {dd}%", f"Importi mensili secondo i millesimi (pagina {i // 11 + 1} di {(len(tab) + 10) // 11}).",
                               riga(0.4 * (len(chunk) + 1) + 0.1, tabella(["Unità", "Proprietario", "Millesimi", "Quota capitale", "Rata/mese", "Recupero fiscale/mese", "Risparmio energ./mese", "IMPATTO REALE/mese"], chunk,
                                                                            [0.7, 1.5, 0.8, 1.1, 1, 1.2, 1.2, 1.2], colore[sc["opzione"]]))))
    # mercato
    m = ctx.get("mercato") or {}
    if m.get("ante") and m.get("post"):
        els = [kpi("VALORE ANTE-INTERVENTO", f"€ {_it(m['ante'], 0)} /m²", "annunci in zona" if m.get("da_annunci") else "stima", "oro"),
               kpi("VALORE POST-INTERVENTO", f"€ {_it(m['post'], 0)} /m²", "immobili efficientati in zona" if m.get("da_annunci") else "stima", "verde"),
               kpi("INCREMENTO POTENZIALE", f"+ € {_it(m['delta'], 0)} /m²", f"{_it(m['delta'] / m['ante'] * 100, 1)}%", "azzurro")]
        mq = ctx.get("mq_esempio") or 100
        S.append(slide("Valorizzazione immobiliare attesa", p.get("zona") or "Confronto di mercato nella zona dell'immobile.", riga(1.3, *els),
                       riga(1.4, card("Come leggere il dato", f"Il confronto usa i prezzi al m² degli immobili in vendita nella zona: così com'è e già efficientati. Un incremento di € {_it(m['delta'], 0)}/m² equivale a circa € {_it(m['delta'] * mq, 0)} su un appartamento di {mq} m², prima di considerare le caratteristiche specifiche dell'unità, piano, esposizione e stato interno.", "azz_c", "azzurro")),
                       riga(0.8, nota("Scenario estimativo FCR basato su annunci pubblici, non perizia.", "grigio_c")),
                       *([riga(0.4 * min(len(ctx.get('annunci') or []), 6) + 0.5, tabella(["Fonte", "Zona", "Stato", "m²", "Prezzo", "€/m²"], [[a.get('Fonte', ''), a.get('Zona', ''), a.get('Stato', ''), _it(a.get('Mq') or 0, 0), _e(a.get('Prezzo') or 0, 0), _e((a.get('Prezzo') or 0) / a['Mq'], 0) if a.get('Mq') else ''] for a in (ctx.get('annunci') or [])[:6]], [1.3, 1.5, 1, .6, 1, 1]))] if ctx.get("annunci") else [])))
        if ctx.get("opzioni") and par and sim and opz:
            o = opz[-1]
            S.append(slide("Costo economico netto e valore patrimoniale", "Lettura congiunta di costo reale e valorizzazione attesa.",
                           riga(1.2, *[kpi(f"OPZ. {oo} - COSTO REALE {int(d)}%", _e(sim[oo][-1][f"costo_reale{int(d)}"] / 1000), f"€/millesimo, {sim[oo][-1]['anni']} anni", colore[oo]) for oo in opz for d in par["detrazioni"][:2]]),
                           riga(1.6, card("Sul piano patrimoniale", f"Lo scenario di valorizzazione di +€ {_it(m['delta'], 0)}/m² va confrontato con il costo economico netto ripartito per millesimi, non con la sola quota capitale e non con il flusso di cassa mensile.", "azz_c", "azzurro"),
                                card("Sul piano economico", "La misura corretta è: capitale finanziato + interessi complessivi − recupero fiscale totale (− risparmio energetico dove dichiarato).", "verde_c", "verde"))))
    # continuita' finanziaria
    cov = ctx.get("copertura") or {}
    tp = cov.get("tipi", {})
    els_cov = []
    if cov and cov.get("totale"):
        els_cov = [kpi("APPROVATO", _e(cov["approvato"], 0), f"{_it(cov['pct_approvato'], 1)}% del progetto", "verde"), kpi("IN ISTRUTTORIA", _e(cov["in_corso"], 0), "", "oro"),
                   kpi("SCOPERTO", _e(cov["scoperto"], 0), "da coprire", "navy")]
    S.append(slide("Continuità finanziaria dell'appalto", "Due finanziamenti paralleli per non fermare il cantiere.",
                   riga(2.2, card("Finanziamento al Condominio", "Richiesto sul codice fiscale del condominio per i lavori sulle parti comuni (copertura, facciate, centrale termica...). Non entra nel merito creditizio del singolo condòmino.", "verde_c", "verde"),
                        card("Finanziamento al condòmino", "Persona fisica, fino al 50% del valore dei lavori sulle parti private. La pratica è inserita direttamente da FCR presso il finanziatore convenzionato (AGOS).", "oro_c", "oro")),
                   *([riga(1.1, *els_cov)] if els_cov else []),
                   riga(0.8, nota("Ogni richiesta è tracciata da FCR (richiesta, istruttoria, esito, erogazione) per garantire la copertura finanziaria dello stato avanzamento lavori."))))
    S.append(slide("Percorso decisionale proposto", "Una sequenza chiara per arrivare all'assemblea con una decisione leggibile.",
                   riga(0.65, passo(1, "Confronto tecnico", "Scelta della soluzione da sottoporre alla delibera.")), riga(0.65, passo(2, "Validazione economica", "Conferma del computo a misura, quadro economico e condizioni di offerta.")),
                   riga(0.65, passo(3, "Richiesta di finanziamento", "Istruttoria preliminare sul Condominio e sui singoli condòmini, senza vincolo verso un istituto specifico.")),
                   riga(0.65, passo(4, "Assemblea", "Presentazione FCR, delibera su soluzione, importo e mandato all'amministratore.", "navy")),
                   riga(0.65, passo(5, "Contratto e cantiere", "Definizione tempi, sicurezza, SAL, contabilità a misura e avvio.", "navy")),
                   riga(0.55, nota("FCR può partecipare all'assemblea per illustrare direttamente differenze tecniche, impatto per millesimo e condizioni operative."))))
    S.append({"tipo": "chiusura", "titolo": "FCR General Contractor", "riga1": "Riqualificare significa proteggere l'edificio, migliorare la prestazione energetica e ricostruire valore.",
              "riga2": f"Validità della proposta: {validita} giorni dalla data di emissione"})
    return S


# ----------------------------------------------------------------------------------------------------------------------------- layout
def _geom(sl):
    """[(riga, [(elemento, x, y, w, h)])] in pollici."""
    out, y, x0, wtot, gap = [], 1.15, 0.6, W - 1.2, 0.18
    for r in sl["righe"]:
        pesi = r["pesi"]
        disp = wtot - gap * (len(r["els"]) - 1)
        x = x0
        celle = []
        for el, pe in zip(r["els"], pesi):
            w = disp * pe / sum(pesi)
            celle.append((el, x, y, w, r["h"]))
            x += w + gap
        out.append(celle)
        y += r["h"] + 0.15
    return out


def _rgb(n):
    c = COL.get(n, n)
    return tuple(int(c[i:i + 2], 16) for i in (0, 2, 4))


# ----------------------------------------------------------------------------------------------------------------------------- PPTX
def pptx_bytes(slides):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    vuota = prs.slide_layouts[6]

    def rgb(n):
        return RGBColor(*_rgb(n))

    def box(s, x, y, w, h, fondo=None, bordo=None, forma=MSO_SHAPE.RECTANGLE):
        sh = s.shapes.add_shape(forma, Inches(x), Inches(y), Inches(w), Inches(h))
        if fondo:
            sh.fill.solid()
            sh.fill.fore_color.rgb = rgb(fondo)
        else:
            sh.fill.background()
        if bordo:
            sh.line.color.rgb = rgb(bordo)
            sh.line.width = Pt(0.75)
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        return sh

    def testo(s, x, y, w, h, t, dim=10, bold=False, colore="testo", align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        for i, linea in enumerate(str(t).split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = linea
            r.font.size, r.font.bold, r.font.name = Pt(dim), bold, "Calibri"
            r.font.color.rgb = rgb(colore)
        return tb

    n = 0
    for sl in slides:
        s = prs.slides.add_slide(vuota)
        n += 1
        if sl["tipo"] in ("copertina", "chiusura"):
            box(s, 0, 0, W, H, "navy")
            box(s, 0.7, 1.4, 8.4, 3.6, "163247")
            if LOGO:
                s.shapes.add_picture(LOGO, Inches(W - 2.1), Inches(0.45), height=Inches(1.4))
            testo(s, 1.0, 1.7, 7.9, 1.6, sl["titolo"], 32, True, "bianco")
            testo(s, 1.0, 3.5, 7.9, 0.5, sl["riga1"], 16, False, "bianco")
            testo(s, 1.0, 4.1, 7.9, 0.5, sl["riga2"], 13, False, "D8E6F0")
            if "data" in sl:
                box(s, W - 4.2, H - 1.7, 3.6, 1.0, "bianco")
                testo(s, W - 4.2, H - 1.65, 3.6, 0.9, sl["data"], 13, True, "titolo", PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
            continue
        box(s, 0, 0, W, H, "sfondo")
        if LOGO:
            s.shapes.add_picture(LOGO, Inches(0.35), Inches(0.2), height=Inches(0.6))
        testo(s, 1.05, 0.15, 11.6, 0.5, sl["titolo"], 24, True, "titolo")
        if sl["sotto"]:
            testo(s, 1.05, 0.62, 11.6, 0.35, sl["sotto"], 11, False, "sotto")
        box(s, 1.05, 1.0, 11.7, 0.015, "bordo")
        for celle in _geom(sl):
            for el, x, y, w, h in celle:
                k = el["k"]
                if k == "kpi":
                    box(s, x, y, w, h, "bianco", "bordo")
                    box(s, x, y, 0.08, h, el["colore"])
                    testo(s, x + 0.15, y + 0.08, w - 0.2, 0.3, el["label"], 8.5, True, "sotto")
                    testo(s, x + 0.15, y + 0.36, w - 0.2, 0.5, el["valore"], 20, True, "titolo")
                    if el["nota"]:
                        testo(s, x + 0.15, y + h - 0.32, w - 0.2, 0.3, el["nota"], 8, False, "sotto")
                elif k == "card":
                    box(s, x, y, w, h, el["fondo"], "bordo")
                    testo(s, x + 0.12, y + 0.08, w - 0.24, 0.35, el["titolo"], 13, True, el["colore"])
                    testo(s, x + 0.12, y + 0.45, w - 0.24, h - 0.5, el["testo"], 10.5)
                elif k == "nota":
                    box(s, x, y, w, h, el["fondo"], "bordo")
                    testo(s, x + 0.1, y, w - 0.2, h, el["testo"], 10, True, "titolo", PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
                elif k == "passo":
                    box(s, x, y, 0.6, h - 0.08, el["colore"])
                    testo(s, x, y, 0.6, h - 0.08, el["n"], 16, True, "bianco", PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
                    testo(s, x + 0.75, y, 3.3, h - 0.08, el["titolo"], 14, True, "titolo", anchor=MSO_ANCHOR.MIDDLE)
                    testo(s, x + 4.1, y, w - 4.1, h - 0.08, el["testo"], 11, anchor=MSO_ANCHOR.MIDDLE)
                elif k == "tab":
                    nc, nr = len(el["colonne"]), len(el["righe"]) + 1
                    rh = min(0.4, (h - 0.05) / nr)
                    tot = sum(el["pesi"])
                    x_ = x
                    for c, pe in zip(el["colonne"], el["pesi"]):
                        cw = w * pe / tot
                        box(s, x_, y, cw, rh, el["colore"])
                        testo(s, x_, y, cw, rh, c, 8.5, True, "bianco", PP_ALIGN.CENTER if c != el["colonne"][0] else PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
                        x_ += cw
                    for i, rr in enumerate(el["righe"]):
                        yy, x_ = y + rh * (i + 1), x
                        ult = el["ultima"] and i == len(el["righe"]) - 1
                        box(s, x, yy, w, rh, "EEF3F7" if ult else ("FFFFFF" if i % 2 == 0 else "F6F8FA"), "bordo")
                        for j, (v, pe) in enumerate(zip(rr, el["pesi"])):
                            cw = w * pe / tot
                            testo(s, x_, yy, cw, rh, v, 9, ult, "titolo" if ult else "testo", PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT if nc <= 3 else PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
                            x_ += cw
        pie = f"FCR General Contractor | Proposta riservata | {sl.get('fonte') or ''}"
        testo(s, 0.4, H - 0.4, 9, 0.3, pie, 8, False, "sotto")
        testo(s, W - 1.0, H - 0.4, 0.6, 0.3, str(n), 10, False, "sotto", PP_ALIGN.RIGHT)
    bio = io.BytesIO()
    prs.save(bio)
    return bio.getvalue()


# ----------------------------------------------------------------------------------------------------------------------------- PDF
def pdf_bytes(slides):
    from fpdf import FPDF
    pdf = FPDF(orientation="L", unit="in", format=(7.5, 13.333))
    pdf.set_auto_page_break(False)
    pdf.add_font("LS", "", FONT_R)
    pdf.add_font("LS", "B", FONT_B)
    pdf.set_title("Proposta FCR")
    pdf.set_author("FCR General Contractor")

    def rect(x, y, w, h, fondo=None, bordo=None):
        if fondo:
            pdf.set_fill_color(*_rgb(fondo))
        if bordo:
            pdf.set_draw_color(*_rgb(bordo))
            pdf.set_line_width(0.01)
        pdf.rect(x, y, w, h, "DF" if fondo and bordo else "F" if fondo else "D" if bordo else "")

    def tx(x, y, w, h, t, dim=10, bold=False, colore="testo", align="L", middle=False):
        pdf.set_font("LS", "B" if bold else "", dim)
        pdf.set_text_color(*_rgb(colore))
        t = str(t)
        lh = dim * 1.25 / 72
        righe_t = []
        for par in t.split("\n"):
            righe_t += pdf.multi_cell(max(w - 0.1, 0.1), lh, par, dry_run=True, output="LINES") or [""]
        yy = y + (max(h - lh * len(righe_t), 0) / 2 if middle else 0.03)
        for ln in righe_t:
            if yy + lh > y + h + 0.02:
                break
            pdf.set_xy(x + 0.05, yy)
            pdf.cell(max(w - 0.1, 0.1), lh, ln, align=align)
            yy += lh

    n = 0
    for sl in slides:
        pdf.add_page()
        n += 1
        if sl["tipo"] in ("copertina", "chiusura"):
            rect(0, 0, W, H, "navy")
            rect(0.7, 1.4, 8.4, 3.6, "163247")
            if LOGO:
                pdf.image(LOGO, W - 2.1, 0.45, h=1.4)
            tx(1.0, 1.7, 7.9, 1.6, sl["titolo"], 32, True, "bianco")
            tx(1.0, 3.5, 7.9, 0.5, sl["riga1"], 16, False, "bianco")
            tx(1.0, 4.1, 7.9, 0.5, sl["riga2"], 13, False, "D8E6F0")
            if "data" in sl:
                rect(W - 4.2, H - 1.7, 3.6, 1.0, "bianco")
                tx(W - 4.2, H - 1.65, 3.6, 0.9, sl["data"], 13, True, "titolo", "C", True)
            continue
        rect(0, 0, W, H, "sfondo")
        if LOGO:
            pdf.image(LOGO, 0.35, 0.2, h=0.6)
        tx(1.05, 0.15, 11.6, 0.5, sl["titolo"], 24, True, "titolo")
        if sl["sotto"]:
            tx(1.05, 0.62, 11.6, 0.35, sl["sotto"], 11, False, "sotto")
        rect(1.05, 1.0, 11.7, 0.015, "bordo")
        for celle in _geom(sl):
            for el, x, y, w, h in celle:
                k = el["k"]
                if k == "kpi":
                    rect(x, y, w, h, "bianco", "bordo")
                    rect(x, y, 0.08, h, el["colore"])
                    tx(x + 0.15, y + 0.08, w - 0.2, 0.3, el["label"], 8.5, True, "sotto")
                    tx(x + 0.15, y + 0.36, w - 0.2, 0.5, el["valore"], 20, True, "titolo")
                    if el["nota"]:
                        tx(x + 0.15, y + h - 0.32, w - 0.2, 0.3, el["nota"], 8, False, "sotto")
                elif k == "card":
                    rect(x, y, w, h, el["fondo"], "bordo")
                    tx(x + 0.12, y + 0.08, w - 0.24, 0.35, el["titolo"], 13, True, el["colore"])
                    tx(x + 0.12, y + 0.45, w - 0.24, h - 0.5, el["testo"], 10.5)
                elif k == "nota":
                    rect(x, y, w, h, el["fondo"], "bordo")
                    tx(x + 0.1, y, w - 0.2, h, el["testo"], 10, True, "titolo", "C", True)
                elif k == "passo":
                    rect(x, y, 0.6, h - 0.08, el["colore"])
                    tx(x, y, 0.6, h - 0.08, el["n"], 16, True, "bianco", "C", True)
                    tx(x + 0.75, y, 3.3, h - 0.08, el["titolo"], 14, True, "titolo", "L", True)
                    tx(x + 4.1, y, w - 4.1, h - 0.08, el["testo"], 11, False, "testo", "L", True)
                elif k == "tab":
                    nc, nr = len(el["colonne"]), len(el["righe"]) + 1
                    rh = min(0.4, (h - 0.05) / nr)
                    tot = sum(el["pesi"])
                    x_ = x
                    for c, pe in zip(el["colonne"], el["pesi"]):
                        cw = w * pe / tot
                        rect(x_, y, cw, rh, el["colore"])
                        tx(x_, y, cw, rh, c, 8.5, True, "bianco", "L" if c == el["colonne"][0] else "C", True)
                        x_ += cw
                    for i, rr in enumerate(el["righe"]):
                        yy, x_ = y + rh * (i + 1), x
                        ult = el["ultima"] and i == len(el["righe"]) - 1
                        rect(x, yy, w, rh, "EEF3F7" if ult else ("FFFFFF" if i % 2 == 0 else "F6F8FA"), "bordo")
                        for j, (v, pe) in enumerate(zip(rr, el["pesi"])):
                            cw = w * pe / tot
                            tx(x_, yy, cw, rh, v, 9, ult, "titolo" if ult else "testo", "L" if j == 0 else "R" if nc <= 3 else "C", True)
                            x_ += cw
        tx(0.4, H - 0.4, 9, 0.3, f"FCR General Contractor | Proposta riservata | {sl.get('fonte') or ''}", 8, False, "sotto")
        tx(W - 1.0, H - 0.4, 0.6, 0.3, str(n), 10, False, "sotto", "R")
    return bytes(pdf.output())
