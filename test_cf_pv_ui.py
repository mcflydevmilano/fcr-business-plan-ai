"""Moduli Cash Flow di cantiere e Computo e Preventivo: UI (AppTest) e report."""
import io

import pytest

from test_hub_bp import _btn, _nuovo


def test_hub_mostra_quattro_strumenti(monkeypatch):
    at = _nuovo(monkeypatch)
    t = " ".join(m.value for m in at.markdown)
    assert "Strumento 3" in t and "Strumento 4" in t
    _btn(at, "Apri Cash Flow")
    _btn(at, "Apri Computo")


def test_apre_cash_flow_e_calcola_l_esempio(monkeypatch):
    at = _nuovo(monkeypatch)
    _btn(at, "Apri Cash Flow").click().run()
    assert not at.exception and len(at.tabs) == 4
    assert any("5.000.000" in m.value for m in at.markdown)
    _btn(at, "Torna all'Hub").click().run()
    assert len(at.tabs) == 0


def test_apre_preventivo(monkeypatch):
    at = _nuovo(monkeypatch)
    _btn(at, "Apri Computo").click().run()
    assert not at.exception and len(at.tabs) == 4


def test_preventivo_invia_al_cash_flow(monkeypatch):
    at = _nuovo(monkeypatch, modulo="pv")
    _btn(at, "Aggiungi al portafoglio").click().run()
    assert not at.exception
    df = at.session_state["cf_df"]
    assert len(df) == 1 and df.iloc[0]["Contratto"] > 0
    _btn(at, "Torna all'Hub").click().run()
    _btn(at, "Apri Cash Flow").click().run()
    assert not at.exception and len(at.session_state["cf_df"]) == 1


def test_parametri_preventivo_sopravvivono_al_cambio_modulo(monkeypatch):
    at = _nuovo(monkeypatch, modulo="pv")
    at.number_input(key="pv_sconto_pct").set_value(4.0).run()
    _btn(at, "Torna all'Hub").click().run()
    _btn(at, "Apri Computo").click().run()
    assert at.number_input(key="pv_sconto_pct").value == 4.0


def test_report_generati(monkeypatch):
    at = _nuovo(monkeypatch, modulo="cf")
    _btn(at, "Genera dossier").click().run()
    assert not at.exception and at.session_state["cf_rep"]["pdf"][:4] == b"%PDF"
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(at.session_state["cf_rep"]["xlsx"]))
    assert {"Input", "Cantieri", "Mensile"} <= set(wb.sheetnames)
    at2 = _nuovo(monkeypatch, modulo="pv")
    _btn(at2, "Genera preventivo").click().run()
    assert not at2.exception and at2.session_state["pv_rep"]["pdf"][:4] == b"%PDF"


def test_pdf_cliente_senza_costi():
    import subprocess, tempfile, os, pv, pv_report
    v, _ = pv.normalizza_voci([{"Descrizione": "Scavo", "Quantita": 10, "Manodopera": 20}])
    o = pv.calcola_preventivo(v, {})
    pdf = pv_report.genera_pdf_pv(pv_report.costruisci_ctx_pv(o, "Cliente", "Lavori"))
    f = os.path.join(tempfile.mkdtemp(), "x.pdf")
    open(f, "wb").write(pdf)
    txt = subprocess.run(["pdftotext", f, "-"], capture_output=True, text=True).stdout.lower()
    assert "costo industriale" not in txt and "utile" not in txt and "spese generali" not in txt


@pytest.mark.parametrize("modulo,chiave,campo,valore", [
    ("cf", "cf_fido", "number_input", 123000.0),
    ("pv", "pv_sconto_pct", "number_input", 4.5),
    ("bp", "bp_ricavi_y1", "number_input", 2_000_000.0),
])
def test_salva_e_riapri_in_ogni_modulo(monkeypatch, modulo, chiave, campo, valore):
    at = _nuovo(monkeypatch, modulo=modulo)
    at.number_input(key=chiave).set_value(valore).run()
    at.text_input(key=f"sv_nome_{modulo}").set_value("Prova").run()
    at.button(key=f"sv_salva_{modulo}").click().run()
    assert not at.exception
    import storage
    assert [p["nome"] for p in storage.elenca_pratiche(modulo=modulo)] == ["Prova"]
    assert storage.elenca_pratiche() == []                       # non compare tra le pratiche del Credit Risk
    at.number_input(key=chiave).set_value(valore + 1).run()
    at.button(key=f"sv_apri_{modulo}").click().run()
    assert not at.exception and at.number_input(key=chiave).value == valore
    at.button(key=f"sv_elimina_{modulo}").click().run()
    assert storage.elenca_pratiche(modulo=modulo) == []


def test_nome_pratica_con_prefisso_riservato(monkeypatch):
    import storage, tempfile, os
    p = os.path.join(tempfile.mkdtemp(), "t.db")
    with pytest.raises(ValueError):
        storage.salva_pratica("cf|x", {}, path=p)
