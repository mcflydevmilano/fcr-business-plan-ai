import math
import os
import shutil
import subprocess
import tempfile

import pandas as pd
import pytest
from openpyxl import load_workbook

import core
import exports


def _ctx(**kw):
    base = dict(
        dati=dict(core.DATI_DEFAULT), leve={"delta_pfn": 50_000, "delta_pn": 100_000, "delta_ebitda": 20_000},
        piano_par=dict(core.PIANO_DEFAULT),
        storico=[{"anno": 2024, "ricavi": 2_400_000.0, "ebitda": 280_000.0, "patrimonio_netto": 700_000.0, "pfn": 650_000.0}],
        commesse_df=core.commesse_vuote(),
        scenari_salvati={"Prudente": {"leve": {"delta_ebitda": -50_000}}},
        dati_demo=True, modello="test-model", pratica="Pratica di prova",
    )
    base.update(kw)
    return exports.costruisci_ctx(**base)


def test_pdf_e_docx_si_generano():
    ctx = _ctx()
    rel = "Solvibilita'.\n\nPunto di forza € 100 – “test” con caratteri speciali."
    pdf = exports.genera_pdf(ctx, rel)
    assert pdf[:4] == b"%PDF" and len(pdf) > 20_000
    docx = exports.genera_docx(ctx, rel)
    assert docx[:2] == b"PK" and len(docx) > 20_000


def test_pdf_senza_commesse_ne_storico():
    ctx = _ctx(commesse_df=None, storico=[], scenari_salvati=None)
    assert exports.genera_pdf(ctx, "ok")[:4] == b"%PDF"
    assert exports.genera_docx(ctx, "ok")[:2] == b"PK"


def test_pdf_con_ebitda_negativo():
    ctx = _ctx(dati=dict(core.DATI_DEFAULT, ebitda=-50_000.0), leve={})
    assert exports.genera_pdf(ctx, "ok")[:4] == b"%PDF"


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice non disponibile")
@pytest.mark.parametrize(
    "dati_extra,leve",
    [
        ({}, {}),
        ({}, {"delta_pfn": 50_000, "delta_pn": 100_000, "delta_ebitda": 20_000}),
        ({"ebitda": -50_000.0}, {}),
        ({"pfn": 3_000_000.0, "patrimonio_netto": 200_000.0}, {"delta_pn": 50_000}),
        ({"pfn": -400_000.0}, {}),
    ],
)
def test_excel_ricalcolato_coincide_con_python(dati_extra, leve):
    dati = dict(core.DATI_DEFAULT, **dati_extra)
    ctx = _ctx(dati=dati, leve=leve, commesse_df=None, storico=[], scenari_salvati=None)
    xlsx = exports.genera_excel(ctx)
    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, "in.xlsx")
        out = os.path.join(tmp, "out")
        os.makedirs(out)
        with open(src, "wb") as f:
            f.write(xlsx)
        subprocess.run(
            ["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, src],
            check=True, capture_output=True, timeout=180,
        )
        wb = load_workbook(os.path.join(out, "in.xlsx"), data_only=True)

    piano, rating = ctx["piano"], ctx["rating"]
    wp, wr = wb["Piano"], wb["Rating"]
    for col, r in zip("CDE", piano):
        assert wp[f"{col}10"].value == pytest.approx(r["Ricavi"])
        assert wp[f"{col}11"].value == pytest.approx(r["EBITDA"])
        assert wp[f"{col}16"].value == pytest.approx(r["FCF"])
        assert wp[f"{col}17"].value == pytest.approx(r["PFN"])
        v = wp[f"{col}18"].value
        if math.isinf(r["PFN/EBITDA"]):
            assert v == "n.s."
        else:
            assert v == pytest.approx(r["PFN/EBITDA"])
        d = wp[f"{col}20"].value
        if r["DSCR"] is None:
            assert d == "n.d."
        else:
            assert d == pytest.approx(r["DSCR"])
    assert wr["B9"].value == rating["punti"]
    assert wr["B10"].value == rating["classe"].split(" - ")[0]
    for col, c in zip("CDE", ctx["cov_res"]):
        assert wp[f"{col}22"].value == c["Esito leva"]
        assert wp[f"{col}23"].value == c["Esito DSCR"]


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice non disponibile")
def test_excel_commesse_formule():
    df = pd.DataFrame([
        {"Commessa": "A", "Cliente": "X", "Importo": 1000.0, "Avanzamento_pct": 50.0, "Fatturato": 400.0, "Incassato": 300.0, "Margine_pct": 10.0},
        {"Commessa": "B", "Cliente": "Y", "Importo": 1000.0, "Avanzamento_pct": 0.0, "Fatturato": 0.0, "Incassato": 0.0, "Margine_pct": 20.0},
    ])
    ctx = _ctx(commesse_df=df, storico=[], scenari_salvati=None)
    kpi = ctx["commesse_kpi"]
    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, "c.xlsx")
        out = os.path.join(tmp, "o")
        os.makedirs(out)
        with open(src, "wb") as f:
            f.write(exports.genera_excel(ctx))
        subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, src],
                       check=True, capture_output=True, timeout=180)
        wc = load_workbook(os.path.join(out, "c.xlsx"), data_only=True)["Commesse"]
    assert wc["H7"].value == pytest.approx(kpi["backlog"])
    assert wc["J7"].value == pytest.approx(kpi["da_fatturare"])
    assert wc["K7"].value == pytest.approx(kpi["crediti"])
    assert wc["H8"].value == pytest.approx(kpi["margine_backlog_pct"])
