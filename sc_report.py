"""sc_report.py - Dossier PDF ed Excel (con formule) dello Scadenzario."""
import versione
import io
import time

import charts
import core
import sc
from core import DISCLAIMER, fmt_num
from pdf_report import BLU, GIALLO_B, GRIGIO, ROSSO_B, VERDE_B, _PDF, _strumenti
from report_util import copertina


def costruisci_ctx_sc(out, docs, nome=""):
    firma = {"k": out["kpi"], "n": nome, "o": str(out["params"]["oggi"])}
    return {
        "sc": out, "docs": docs, "nome": nome or "", "modello": "nessuno", "pratica": nome or "",
        "report_id": core.hash_payload(firma, 10).upper(), "generato_il": time.strftime("%d/%m/%Y %H:%M"),
        "header_titolo": "FCR GENERAL CONTRACTOR - Scadenzario", "header_sub": "Incassi e pagamenti | " + (nome or "scadenzario"),
        "riga_cover": ("Scadenzario: " + nome) if nome else "Scadenzario incassi e pagamenti",
        "titolo_documento": "FCR General Contractor - Scadenzario", "footer_dx": "Powered by Aetheris Labs",
        "nota_cover": "Previsione basata sulle scadenze inserite e su ipotesi dell'utente: non e' una garanzia di incasso o di pagamento.",
    }


def genera_pdf_sc(ctx):
    out, docs = ctx["sc"], ctx["docs"]
    q, k, sett = out["params"], out["kpi"], out["settimane"]
    pdf = _PDF(ctx)
    grafici = {"sett": charts.fig_png(charts.fig_sc_settimane(sett, q["fido"] or None, "light")),
               "fasce": charts.fig_png(charts.fig_sc_fasce(out["fasce"], "light"))}
    T = _strumenti(pdf, grafici)
    n = [0]

    def sez(t, r=80):
        n[0] += 1
        T.sezione(t, n[0], r)

    if k["prima_settimana_fuori_fido"]:
        esito = ("SCOPERTO OLTRE IL FIDO", ROSSO_B, f"Dalla settimana {k['prima_settimana_fuori_fido']}: picco {fmt_num(k['picco_fabbisogno'])} EUR contro un fido di {fmt_num(k['fido'])} EUR")
    elif not k["ok_settimana"]:
        esito = ("CASSA NEGATIVA QUESTA SETTIMANA", GIALLO_B, f"Fine settimana 1: {fmt_num(k['cassa_settimana1'])} EUR (dentro il fido di {fmt_num(k['fido'])} EUR)")
    elif k["picco_fabbisogno"] > 0:
        esito = ("SCOPERTO DENTRO IL FIDO", GIALLO_B, f"Picco di {fmt_num(k['picco_fabbisogno'])} EUR alla settimana {k['settimana_picco']} su un fido di {fmt_num(k['fido'])} EUR")
    else:
        esito = ("CASSA SOSTENIBILE", VERDE_B, f"Fine settimana 1: {fmt_num(k['cassa_settimana1'])} EUR; minimo previsto {fmt_num(k['cassa_minima'])} EUR")
    copertina(pdf, T.tile, ctx, "Scadenzario", "Incassi e pagamenti", f"Situazione al {q['oggi'].strftime('%d/%m/%Y')}", [
        ("DATA", ctx["generato_il"]), ("ID REPORT", ctx["report_id"]), ("DOCUMENTI", str(k["n_documenti"])), ("ORIZZONTE", f"{k['settimane']} settimane")],
        [("Da incassare", fmt_num(k["da_incassare"]) + " EUR", "residuo clienti", BLU),
         ("Da pagare", fmt_num(k["da_pagare"]) + " EUR", "residuo fornitori", BLU),
         ("Crediti scaduti", fmt_num(k["scaduto_attivo"]) + " EUR", f"{fmt_num(k['scaduto_attivo_pct'], 0)}% dei crediti", ROSSO_B if k["scaduto_attivo_pct"] > 25 else GIALLO_B if k["scaduto_attivo_pct"] > 10 else VERDE_B),
         ("Cassa fine sett. 1", fmt_num(k["cassa_settimana1"]) + " EUR", "dopo 7 giorni", VERDE_B if k["ok_settimana"] else ROSSO_B)], esito)
    pdf.add_page()
    sez("Previsione settimana per settimana", 100)
    righe = [(str(r["Settimana"]), r["Dal"].strftime("%d/%m") + "-" + r["Al"].strftime("%d/%m"), fmt_num(r["Incassi"]), fmt_num(-r["Pagamenti"]),
              fmt_num(r["Flusso netto"]), fmt_num(r["Cassa finale"]), fmt_num(r["Scoperto"])) for r in sett]
    righe.append(("Tot.", "", fmt_num(k["incassi_tot"]), fmt_num(-k["pagamenti_tot"]), fmt_num(k["incassi_tot"] - k["pagamenti_tot"]), fmt_num(k["cassa_finale"]), ""))
    T.tabella(["Sett.", "Periodo", "Incassi", "Pagamenti", "Flusso", "Cassa", "Scoperto"], righe, [14, 28, 28, 28, 28, 32, 24], ["C", "C", "R", "R", "R", "R", "R"], size=8, h=5.6)
    T.immagine("sett", 160)
    sez("Anzianita' dei documenti", 90)
    T.immagine("fasce", 140)
    T.tabella(["Fascia", "Da incassare", "Da pagare"], [(f["Fascia"], fmt_num(f["Attive"]), fmt_num(f["Passive"])) for f in out["fasce"]], [70, 50, 50], ["L", "R", "R"], size=8.5)
    sol = sc.solleciti(docs, q["oggi"])
    sez("Clienti da sollecitare", 60)
    if sol:
        T.tabella(["Cliente", "Scaduto", "Ritardo max", "Livello"], [(c["Cliente"][:30], fmt_num(c["Totale scaduto"]), f"{c['Ritardo max (gg)']} gg", c["Livello"]) for c in sol[:20]],
                  [66, 34, 30, 52], ["L", "R", "C", "L"], size=8.5)
    else:
        T.paragrafo("Nessun credito scaduto da sollecitare.", 9)
    sez("Punti di attenzione", 40)
    for a in out["avvisi"] or ["I controlli automatici non hanno rilevato criticita'."]:
        T.paragrafo("- " + a, 9)
    sez("Ipotesi e limiti", 70)
    T.paragrafo(f"Fatture attive non scadute: incasso alla scadenza + {q['ritardo_attivo_gg']} giorni di ritardo abituale. Crediti scaduti: il {fmt_num(q['incasso_scaduti_pct'], 0)}% si assume incassato "
                f"nella settimana 1, il resto e' considerato incerto e non entra nella previsione. Debiti scaduti: il {fmt_num(q['pagamento_scaduti_pct'], 0)}% si assume pagato nella settimana 1. "
                f"Saldo di conto iniziale {fmt_num(q['cassa0'])} EUR, fido {fmt_num(q['fido'])} EUR. Gli importi sono IVA inclusa, cioe' quelli che movimentano il conto.", 8.5)
    T.paragrafo("Limiti: la previsione riflette solo le fatture inserite (non stipendi, imposte, rate di finanziamento, costi non ancora fatturati) e le date di scadenza dichiarate.", 8.5)
    pdf.ln(4)
    pdf.set_font("Helvetica", "I", 8)
    pdf.set_text_color(*GRIGIO)
    pdf.set_x(14)
    pdf.multi_cell(182, 4.1, "Avvertenza: " + DISCLAIMER.replace("modello di scoring interno", "modello di previsione").encode("latin-1", "replace").decode("latin-1"))
    return bytes(pdf.output())


def genera_excel_sc(ctx):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    out, docs = ctx["sc"], ctx["docs"]
    q, sett = out["params"], out["settimane"]
    wb = Workbook()
    giallo, testata = PatternFill("solid", fgColor="FFF2CC"), PatternFill("solid", fgColor="1E293B")
    f_h, f_in = Font(color="FFFFFF", bold=True), Font(color="1D4ED8", bold=True)
    EUR = '#,##0;[Red]-#,##0'
    ws = wb.active
    ws.title = "Parametri"
    ws["A1"] = "Scadenzario - parametri (celle gialle modificabili)"
    ws["A1"].font = Font(bold=True, size=13)
    for i, (e, v, fm) in enumerate([("Data di riferimento", q["oggi"], "DD/MM/YYYY"), ("Saldo di conto (EUR)", q["cassa0"], EUR), ("Fido disponibile (EUR)", q["fido"], EUR)], start=3):
        ws.cell(row=i, column=1, value=e)
        c = ws.cell(row=i, column=2, value=v)
        c.fill, c.font, c.number_format = giallo, f_in, fm
    ws["A7"] = "Nota: i giorni di ritardo e i residui si ricalcolano con la data di riferimento; la previsione settimanale e' quella calcolata dall'app."
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18
    wd = wb.create_sheet("Documenti")
    for j, t in enumerate(["Tipo", "Controparte", "Cantiere", "Documento", "Data documento", "Scadenza", "Importo", "Saldato", "Residuo", "Giorni di ritardo"], 1):
        c = wd.cell(row=1, column=j, value=t)
        c.fill, c.font = testata, f_h
    for i, d in enumerate(docs, start=2):
        for j, key in enumerate(["Tipo", "Controparte", "Cantiere", "Documento"], 1):
            wd.cell(row=i, column=j, value=d[key])
        wd.cell(row=i, column=5, value=sc.data_da(d["Data_doc"])).number_format = "DD/MM/YYYY"
        wd.cell(row=i, column=6, value=sc.data_da(d["Scadenza"])).number_format = "DD/MM/YYYY"
        wd.cell(row=i, column=7, value=d["Importo"]).number_format = EUR
        wd.cell(row=i, column=8, value=d["Saldato"]).number_format = EUR
        wd.cell(row=i, column=9, value=f"=G{i}-H{i}").number_format = EUR
        wd.cell(row=i, column=10, value=f"=IF(I{i}>0,MAX(0,Parametri!$B$3-F{i}),0)")
    u = len(docs) + 1
    wd.cell(row=u + 2, column=1, value="Da incassare").font = Font(bold=True)
    wd.cell(row=u + 2, column=9, value=f'=SUMIFS(I2:I{u},A2:A{u},"Attiva")').number_format = EUR
    wd.cell(row=u + 3, column=1, value="Da pagare").font = Font(bold=True)
    wd.cell(row=u + 3, column=9, value=f'=SUMIFS(I2:I{u},A2:A{u},"Passiva")').number_format = EUR
    wd.cell(row=u + 4, column=1, value="Crediti scaduti").font = Font(bold=True)
    wd.cell(row=u + 4, column=9, value=f'=SUMIFS(I2:I{u},A2:A{u},"Attiva",J2:J{u},">0")').number_format = EUR
    for col, w in zip("ABCDEFGHIJ", [10, 28, 22, 14, 16, 14, 14, 14, 14, 16]):
        wd.column_dimensions[col].width = w
    wd.freeze_panes = "A2"
    wsx = wb.create_sheet("Settimane")
    for j, t in enumerate(["Settimana", "Dal", "Al", "Incassi", "Pagamenti", "Flusso netto", "Cassa finale", "Scoperto", "Fido residuo"], 1):
        c = wsx.cell(row=1, column=j, value=t)
        c.fill, c.font = testata, f_h
    for i, r in enumerate(sett, start=2):
        prec = "Parametri!$B$4" if i == 2 else f"G{i - 1}"
        wsx.cell(row=i, column=1, value=r["Settimana"])
        wsx.cell(row=i, column=2, value=r["Dal"]).number_format = "DD/MM/YYYY"
        wsx.cell(row=i, column=3, value=r["Al"]).number_format = "DD/MM/YYYY"
        wsx.cell(row=i, column=4, value=round(r["Incassi"], 2)).number_format = EUR
        wsx.cell(row=i, column=5, value=round(r["Pagamenti"], 2)).number_format = EUR
        wsx.cell(row=i, column=6, value=f"=D{i}-E{i}").number_format = EUR
        wsx.cell(row=i, column=7, value=f"={prec}+F{i}").number_format = EUR
        wsx.cell(row=i, column=8, value=f"=MAX(0,-G{i})").number_format = EUR
        wsx.cell(row=i, column=9, value=f"=Parametri!$B$5-H{i}").number_format = EUR
    n = len(sett) + 1
    wsx.cell(row=n + 2, column=1, value="Cassa minima").font = Font(bold=True)
    wsx.cell(row=n + 2, column=4, value=f"=MIN(G2:G{n})").number_format = EUR
    wsx.cell(row=n + 3, column=1, value="Picco di scoperto").font = Font(bold=True)
    wsx.cell(row=n + 3, column=4, value=f"=MAX(H2:H{n})").number_format = EUR
    for col in "ABCDEFGHI":
        wsx.column_dimensions[col].width = 15
    wsx.freeze_panes = "B2"
    versione.aggiungi_foglio_info(wb, "Scadenzario", __import__("time").strftime("%d/%m/%Y %H:%M"))
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
