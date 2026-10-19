"""sl_report.py - Dossier PDF ed Excel (con formule) di SAL e contabilita' lavori."""
import io
import time

import charts
import core
from core import DISCLAIMER, fmt_num
from pdf_report import BLU, GIALLO_B, GRIGIO, ROSSO_B, VERDE_B, _PDF, _strumenti
from report_util import copertina
from sc import data_da

MAX_COMMESSE_DETTAGLIO = 6


def costruisci_ctx_sl(out, commesse, sal, costi, nome=""):
    firma = {"k": out["kpi"], "n": nome, "o": str(out["params"]["oggi"])}
    return {"sl": out, "commesse_in": commesse, "sal_in": sal, "costi_in": costi, "nome": nome or "", "modello": "nessuno", "pratica": nome or "",
            "report_id": core.hash_payload(firma, 10).upper(), "generato_il": time.strftime("%d/%m/%Y %H:%M"),
            "header_titolo": "FCR GENERAL CONTRACTOR - SAL e contabilita' lavori", "header_sub": "Controllo commesse | " + (nome or "portafoglio lavori"),
            "titolo_documento": "FCR General Contractor - SAL e contabilita' lavori", "footer_dx": "Powered by Aetheris Labs"}


def _fd(s):
    d = data_da(s)
    return d.strftime("%d/%m/%Y") if d else ""


def genera_pdf_sl(ctx):
    out = ctx["sl"]
    k, q, det = out["kpi"], out["params"], out["commesse"]
    pdf = _PDF(ctx)
    grafici = {"margini": charts.fig_png(charts.fig_sl_margini(det, "light"))}
    for d in sorted(det, key=lambda x: x["Margine finale %"])[:MAX_COMMESSE_DETTAGLIO]:
        grafici["c_" + d["Commessa"]] = charts.fig_png(charts.fig_sl_commessa(d, "light"))
    T = _strumenti(pdf, grafici)
    n = [0]

    def sez(t, r=80):
        n[0] += 1
        T.sezione(t, n[0], r)

    if k["n_bad"]:
        esito = ("COMMESSE IN SOFFERENZA", ROSSO_B, f"{k['n_bad']} commesse con margine finale a rischio su {k['n_commesse']}")
    elif k["n_warn"]:
        esito = ("COMMESSE DA TENERE D'OCCHIO", GIALLO_B, f"{k['n_warn']} commesse con scostamenti su {k['n_commesse']}")
    else:
        esito = ("COMMESSE IN LINEA", VERDE_B, f"{k['n_commesse']} commesse in linea con il preventivo")
    copertina(pdf, T.tile, ctx, "SAL e contabilita' lavori", "Controllo delle commesse", f"Situazione al {q['oggi'].strftime('%d/%m/%Y')}", [
        ("DATA", ctx["generato_il"]), ("ID REPORT", ctx["report_id"]), ("COMMESSE", str(k["n_commesse"])), ("SAL EMESSI", str(k["n_sal"]))],
        [("Contratti", fmt_num(k["contratti"]) + " EUR", f"avanzamento {fmt_num(k['avanzamento'], 0)}%", BLU),
         ("Margine previsto", f"{fmt_num(k['margine_finale_pct'], 1)}%", f"preventivato {fmt_num(k['margine_atteso_pct'], 1)}%", VERDE_B if k["margine_finale_pct"] >= k["margine_atteso_pct"] - 1 else ROSSO_B),
         ("Da fatturare", fmt_num(k["da_fatturare"]) + " EUR", "lavori eseguiti", GIALLO_B if k["da_fatturare"] > 0 else VERDE_B),
         ("Crediti scaduti", fmt_num(k["crediti_scaduti"]) + " EUR", "SAL non incassati", ROSSO_B if k["crediti_scaduti"] > 0 else VERDE_B)], esito)
    pdf.add_page()
    sez("Quadro delle commesse", 100)
    righe = [(d["Commessa"][:22], fmt_num(d["Contratto totale"]), fmt_num(d["Avanzamento %"], 0) + "%", fmt_num(d["Costi sostenuti"]), fmt_num(d["Margine atteso %"], 1) + "%",
              fmt_num(d["Margine finale %"], 1) + "%", {"ok": "OK", "warn": "ATTENZ.", "bad": "CRITICA"}[d["Stato"]]) for d in det]
    T.tabella(["Commessa", "Contratto", "Avanz.", "Costi sost.", "Marg. prev.", "Marg. finale", "Stato"], righe, [42, 30, 16, 30, 24, 26, 20], ["L", "R", "C", "R", "R", "R", "C"], size=8, h=5.8)
    T.immagine("margini", 160)
    sez("Fatturazione e incassi", 70)
    T.tabella(["Voce", "Importo (EUR)"], [
        ("Lavori eseguiti", fmt_num(k["eseguito"])), ("Fatturato (imponibile + ritenute)", fmt_num(k["fatturato"])), ("Ancora da fatturare", fmt_num(k["da_fatturare"])),
        ("Incassato (IVA inclusa)", fmt_num(k["incassato"])), ("Crediti scaduti", fmt_num(k["crediti_scaduti"])), ("Ritenute di garanzia trattenute", fmt_num(k["ritenute"])),
        ("Costi sostenuti", fmt_num(k["costi_sostenuti"])), ("Debiti verso fornitori", fmt_num(k["debiti_fornitori"])), ("di cui scaduti", fmt_num(k["fornitori_scaduti"]))],
        [120, 62], ["L", "R"])
    pdf.add_page()
    sez("Commesse da guardare da vicino", 60)
    for d in sorted(det, key=lambda x: x["Margine finale %"])[:MAX_COMMESSE_DETTAGLIO]:
        if pdf.get_y() > 200:
            pdf.add_page()
        T.sottotitolo(f"{d['Commessa']}" + (f" - {d['Cliente']}" if d["Cliente"] else ""))
        T.immagine("c_" + d["Commessa"], 120)
        cpi = "n.d." if d["CPI"] is None else fmt_num(d["CPI"], 2)
        T.paragrafo(f"Avanzamento {fmt_num(d['Avanzamento %'], 1)}% (fatturato {fmt_num(d['Avanzamento fatturato %'], 1)}%) | indice di costo CPI {cpi} | costi a finire "
                    f"{fmt_num(d['Costi a finire (EAC)'])} EUR | margine finale previsto {fmt_num(d['Margine finale previsto'])} EUR ({fmt_num(d['Margine finale %'], 1)}%) "
                    f"contro {fmt_num(d['Margine atteso %'], 1)}% preventivato.", 8.5)
        sal = [f for f in out["fatture"] if f["Commessa"] == d["Commessa"]]
        if sal:
            T.tabella(["SAL", "Data", "Avanz.", "Imponibile", "Totale", "Incassato", "Residuo"],
                      [(f"{f['Tipo']} {f['Numero']}".strip(), _fd(f["Data"]), fmt_num(f["Avanzamento"], 0) + "%", fmt_num(f["Imponibile"]), fmt_num(f["Totale"]), fmt_num(f["Incassato"]),
                        fmt_num(f["Residuo"]) + (" !" if f["Scaduta"] else "")) for f in sal],
                      [30, 24, 18, 28, 28, 28, 26], ["L", "C", "C", "R", "R", "R", "R"], size=7.8, h=5.4)
    sez("Punti di attenzione", 40)
    for a in out["avvisi"] or ["I controlli automatici non hanno rilevato criticita'."]:
        T.paragrafo("- " + a, 9)
    sez("Metodo e limiti", 70)
    T.paragrafo("Ogni SAL fattura la quota di contratto maturata dall'ultimo avanzamento, al netto del recupero dell'anticipo e della ritenuta di garanzia, piu' IVA. "
                "Il margine a finire usa il metodo del valore maturato: valore maturato (EV) = costi preventivati x avanzamento; indice di costo CPI = EV / costi sostenuti; "
                "costi a finire attesi = costi preventivati / CPI (con avanzamento sotto il 5% si usa il preventivo). La stima presume che l'efficienza dei costi osservata continui.", 8.5)
    T.paragrafo("Limiti: l'avanzamento reale e' un dato inserito dall'utente; non sono modellati riserve, penali, revisione prezzi, subappalti in contabilita' separata e le "
                "ritenute sono considerate al netto di IVA fino allo svincolo.", 8.5)
    pdf.ln(4)
    pdf.set_font("Helvetica", "I", 8)
    pdf.set_text_color(*GRIGIO)
    pdf.set_x(14)
    pdf.multi_cell(182, 4.1, "Avvertenza: " + DISCLAIMER.replace("modello di scoring interno", "modello di controllo").encode("latin-1", "replace").decode("latin-1"))
    return bytes(pdf.output())


def genera_excel_sl(ctx):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    out = ctx["sl"]
    det = out["commesse"]
    wb = Workbook()
    testata, giallo = PatternFill("solid", fgColor="1E293B"), PatternFill("solid", fgColor="FFF2CC")
    f_h, f_in = Font(color="FFFFFF", bold=True), Font(color="1D4ED8", bold=True)
    EUR = '#,##0;[Red]-#,##0'

    def intest(ws, cols):
        for j, t in enumerate(cols, 1):
            c = ws.cell(row=1, column=j, value=t)
            c.fill, c.font = testata, f_h
        ws.freeze_panes = "A2"

    wc = wb.active
    wc.title = "Commesse"
    intest(wc, ["Commessa", "Contratto", "Varianti", "Contratto totale", "Costi preventivati", "Avanzamento %", "Costi sostenuti", "Valore maturato (EV)", "CPI",
                "Costi a finire (EAC)", "Margine preventivato", "Margine prev. %", "Margine finale", "Margine finale %", "Scostamento (pp)"])
    ci = {c["Commessa"]: c for c in ctx["commesse_in"]}
    for i, d in enumerate(det, start=2):
        c = ci.get(d["Commessa"], {})
        wc.cell(row=i, column=1, value=d["Commessa"])
        for col, v in ((2, c.get("Contratto", 0)), (3, c.get("Varianti", 0)), (5, d["Costi preventivati"]), (6, d["Avanzamento %"])):
            cell = wc.cell(row=i, column=col, value=v)
            cell.fill, cell.font = giallo, f_in
            cell.number_format = EUR if col != 6 else "0.0"
        wc.cell(row=i, column=4, value=f"=B{i}+C{i}").number_format = EUR
        wc.cell(row=i, column=7, value=f'=SUMIFS(Costi!$F:$F,Costi!$A:$A,A{i})').number_format = EUR
        wc.cell(row=i, column=8, value=f"=E{i}*F{i}/100").number_format = EUR
        wc.cell(row=i, column=9, value=f"=IF(AND(G{i}>0,H{i}>0),H{i}/G{i},\"\")").number_format = "0.00"
        wc.cell(row=i, column=10, value=f"=MAX(G{i},IF(AND(F{i}>=5,G{i}>0,H{i}>0),E{i}/(H{i}/G{i}),E{i}))").number_format = EUR
        wc.cell(row=i, column=11, value=f"=D{i}-E{i}").number_format = EUR
        wc.cell(row=i, column=12, value=f"=IF(D{i}=0,0,K{i}/D{i}*100)").number_format = "0.0"
        wc.cell(row=i, column=13, value=f"=D{i}-J{i}").number_format = EUR
        wc.cell(row=i, column=14, value=f"=IF(D{i}=0,0,M{i}/D{i}*100)").number_format = "0.0"
        wc.cell(row=i, column=15, value=f"=N{i}-L{i}").number_format = "0.0"
    u = len(det) + 1
    wc.cell(row=u + 1, column=1, value="Totale").font = Font(bold=True)
    for col in "BCDEGHJKM":
        wc[f"{col}{u + 1}"] = f"=SUM({col}2:{col}{u})"
        wc[f"{col}{u + 1}"].number_format = EUR
    wc[f"L{u + 1}"] = f"=IF(D{u + 1}=0,0,K{u + 1}/D{u + 1}*100)"
    wc[f"N{u + 1}"] = f"=IF(D{u + 1}=0,0,M{u + 1}/D{u + 1}*100)"
    wc.column_dimensions["A"].width = 26
    for col in "BCDEFGHIJKLMNO":
        wc.column_dimensions[col].width = 16
    ws = wb.create_sheet("SAL")
    intest(ws, ["Commessa", "N. SAL", "Tipo", "Data", "Avanzamento %", "Lordo", "Recupero anticipo", "Ritenuta", "Imponibile", "IVA", "Totale", "Scadenza", "Incassato", "Residuo"])
    for i, f in enumerate(out["fatture"], start=2):
        for j, v in enumerate([f["Commessa"], f["Numero"], f["Tipo"], data_da(f["Data"]), f["Avanzamento"], f["Lordo"], f["Recupero anticipo"], f["Ritenuta"]], 1):
            cell = ws.cell(row=i, column=j, value=v)
            if j == 4:
                cell.number_format = "DD/MM/YYYY"
            elif j >= 6:
                cell.number_format = EUR
        ws.cell(row=i, column=9, value=f"=F{i}-G{i}-H{i}").number_format = EUR
        ws.cell(row=i, column=10, value=round(f["IVA"], 2)).number_format = EUR
        ws.cell(row=i, column=11, value=f"=I{i}+J{i}").number_format = EUR
        ws.cell(row=i, column=12, value=data_da(f["Scadenza"])).number_format = "DD/MM/YYYY"
        ws.cell(row=i, column=13, value=f["Incassato"]).number_format = EUR
        ws.cell(row=i, column=14, value=f"=K{i}-M{i}").number_format = EUR
    for col, w in zip("ABCDEFGHIJKLMN", [26, 10, 10, 13, 14, 14, 16, 12, 14, 12, 14, 13, 14, 14]):
        ws.column_dimensions[col].width = w
    wk = wb.create_sheet("Costi")
    intest(wk, ["Commessa", "Data", "Fornitore", "Descrizione", "Categoria", "Importo", "Pagato", "Da pagare"])
    for i, x in enumerate(ctx["costi_in"], start=2):
        for j, v in enumerate([x["Commessa"], data_da(x["Data"]), x["Fornitore"], x["Descrizione"], x["Categoria"], x["Importo"], x["Pagato"]], 1):
            cell = wk.cell(row=i, column=j, value=v)
            if j == 2:
                cell.number_format = "DD/MM/YYYY"
            elif j >= 6:
                cell.number_format = EUR
        wk.cell(row=i, column=8, value=f"=F{i}-G{i}").number_format = EUR
    for col, w in zip("ABCDEFGH", [26, 13, 26, 34, 16, 14, 14, 14]):
        wk.column_dimensions[col].width = w
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
